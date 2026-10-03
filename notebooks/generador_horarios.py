#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
SISTEMA DE PROGRAMACIÓN AUTOMATIZADA DE HORARIOS ESCOLARES — AÑO ESCOLAR 2026
Colegio Madres Dominicas (MMDD) - Concepción, Chile
Taller de Gestión de Operaciones (TGOP) - Ingeniería Civil Industrial, UdeC
===============================================================================

Este script implementa un motor algorítmico capaz de construir mallas horarias
factibles para los 24 cursos regulares del establecimiento (1° Básico a 4° Medio,
secciones A y B) y para la Educación Física de los 3 niveles de párvulos, usando
los parámetros oficiales del año 2026 (carpeta data/Data 2026):
  - HORARIO CURSOS 2026.xlsx       -> bloques, jornadas, cargas, electivos y recintos.
  - DISTRIBUCIÓN HORARIA 2026.docx -> dotación docente y asignación profesor-curso.

Módulos incluidos:
1. DatosColegio: bloques, jornadas, dotación docente, mallas, electivos y recintos 2026.
2. HorarioEscolar: Estructura de datos matricial para asignaciones y consultas.
3. ValidadorRestricciones: Auditoría exhaustiva de restricciones duras y blandas.
4. MotorHorarios: Heurística constructiva + búsqueda local Min-Conflicts con lista tabú.
5. ExportadorExcel: Libros .xlsx con el mismo diseño que 'HORARIO CURSOS 2026.xlsx' (openpyxl).
6. MenuInteractivo: Interfaz de consola para consultar cursos, docentes y exportar.
"""

import io
import os
import sys
import random
import zipfile
import argparse
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.drawing.image import Image as ImagenExcel
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.indexed_list import IndexedList
    from openpyxl.worksheet.page import PageMargins
except ImportError:  # El motor y la consola funcionan sin openpyxl; solo la exportación lo requiere
    openpyxl = None

# =============================================================================
# 1. PARÁMETROS Y DATOS DEL COLEGIO MADRES DOMINICAS (AÑO 2026)
# =============================================================================

class DatosColegio:
    """Catálogo y parámetros institucionales del Colegio Madres Dominicas (año escolar 2026)."""

    ANIO = 2026
    DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes']

    # Bloques pedagógicos de 45 minutos (fuente: HORARIO CURSOS 2026.xlsx)
    HORARIOS_BLOQUES = {
        1: ('08:00', '08:45'),
        2: ('08:45', '09:30'),
        3: ('09:45', '10:30'),
        4: ('10:30', '11:15'),
        5: ('11:30', '12:15'),
        6: ('12:15', '13:00'),
        7: ('13:10', '13:55'),
        8: ('13:55', '14:40'),
        9: ('15:10', '15:55'),   # Jornada tarde: solo lunes en 3° y 4° medio
        10: ('15:55', '16:40')   # Jornada tarde: solo lunes en 3° y 4° medio
    }

    # Pausa que sigue a cada bloque: (nombre, rango horario)
    RECREOS = {
        2: ('RECREO', '09:30 - 09:45'),
        4: ('RECREO', '11:15 - 11:30'),
        6: ('RECREO', '13:00 - 13:10'),
        8: ('COLACIÓN', '14:40 - 15:10')
    }

    # Pares pedagógicos de 90 minutos (un bloque doble nunca cruza un recreo)
    PARES = [(1, 2), (3, 4), (5, 6), (7, 8), (9, 10)]

    # Los 24 cursos regulares del colegio
    CURSOS = [
        '1° BÁSICO A', '1° BÁSICO B',
        '2° BÁSICO A', '2° BÁSICO B',
        '3° BÁSICO A', '3° BÁSICO B',
        '4° BÁSICO A', '4° BÁSICO B',
        '5° BÁSICO A', '5° BÁSICO B',
        '6° BÁSICO A', '6° BÁSICO B',
        '7° BÁSICO A', '7° BÁSICO B',
        '8° BÁSICO A', '8° BÁSICO B',
        '1° MEDIO A', '1° MEDIO B',
        '2° MEDIO A', '2° MEDIO B',
        '3° MEDIO A', '3° MEDIO B',
        '4° MEDIO A', '4° MEDIO B'
    ]

    # Párvulos: solo se programa su Educación Física (comparten docente y Gimnasio B con básica)
    CURSOS_PARVULOS = ['PREKINDER A', 'KINDER A', 'KINDER B']

    # Último bloque lectivo de cada día (lunes a viernes) según las mallas 2026.
    # En los datos oficiales 5°A/5°B y 8°A/8°B no comparten la misma jornada.
    JORNADAS = {
        '1° BÁSICO A': (8, 8, 7, 8, 7), '1° BÁSICO B': (8, 8, 7, 8, 7),
        '2° BÁSICO A': (8, 8, 7, 8, 7), '2° BÁSICO B': (8, 8, 7, 8, 7),
        '3° BÁSICO A': (8, 8, 7, 8, 7), '3° BÁSICO B': (8, 8, 7, 8, 7),
        '4° BÁSICO A': (8, 8, 7, 8, 7), '4° BÁSICO B': (8, 8, 7, 8, 7),
        '5° BÁSICO A': (8, 8, 8, 6, 8), '5° BÁSICO B': (8, 8, 8, 8, 6),
        '6° BÁSICO A': (8, 8, 7, 8, 6), '6° BÁSICO B': (8, 8, 7, 8, 6),
        '7° BÁSICO A': (8, 8, 7, 8, 6), '7° BÁSICO B': (8, 8, 7, 8, 6),
        '8° BÁSICO A': (6, 8, 7, 8, 8), '8° BÁSICO B': (8, 8, 7, 8, 6),
        '1° MEDIO A': (8, 8, 8, 8, 8), '1° MEDIO B': (8, 8, 8, 8, 8),
        '2° MEDIO A': (8, 8, 8, 8, 8), '2° MEDIO B': (8, 8, 8, 8, 8),
        '3° MEDIO A': (10, 8, 8, 8, 8), '3° MEDIO B': (10, 8, 8, 8, 8),
        '4° MEDIO A': (10, 8, 8, 8, 8), '4° MEDIO B': (10, 8, 8, 8, 8),
        # Supuesto: la Ed. Física de párvulos se dicta en la jornada de mañana (bloques 1 a 6)
        'PREKINDER A': (6, 6, 6, 6, 6), 'KINDER A': (6, 6, 6, 6, 6), 'KINDER B': (6, 6, 6, 6, 6),
    }

    # Dotación docente 2026, agrupada como en DISTRIBUCIÓN HORARIA 2026.docx
    DEPARTAMENTOS = [
        ('LENGUAJE – FILOSOFÍA - ASIGNATURAS DE PROFUNDIZACIÓN',
         ['Lenguaje 1', 'Lenguaje 2', 'Lenguaje 3', 'Lenguaje 4', 'Lenguaje 5']),
        ('MATEMÁTICA - ASIGNATURAS DE PROFUNDIZACIÓN',
         ['Matemática 1', 'Matemática 2', 'Matemática 3', 'Matemática 4']),
        ('INGLÉS',
         ['Inglés 1', 'Inglés 2', 'Inglés 3', 'Inglés 4', 'Inglés 5']),
        ('HISTORIA, GEOGRAFÍA Y C. SOCIALES – EDUCACIÓN CIUDADANA - ASIGNATURAS DE PROFUNDIZACIÓN',
         ['Historia 1', 'Historia 2', 'Historia 3']),
        ('CIENCIAS NATURALES – CIENCIAS PARA LA CIUDADANÍA - ASIGNATURAS DE PROFUNDIZACIÓN – RELIGIÓN',
         ['Química', 'Biología', 'C. Naturales y Religión', 'C. Naturales', 'Física', 'Religión']),
        ('ARTES VISUALES – MÚSICA – TECNOLOGÍA - ASIGNATURAS DE PROFUNDIZACIÓN',
         ['Artes y Tecnología 1', 'Artes y Tecnología 2', 'Música']),
        ('EDUCACIÓN FÍSICA',
         ['Ed. Física 1', 'Ed. Física 2', 'Ed. Física 3']),
        ('EDUCACIÓN GENERAL BÁSICA',
         ['Básica 1', 'Básica 2', 'Básica 3', 'Básica 4', 'Básica 5', 'Básica 6', 'Básica 7', 'Básica 8']),
    ]

    # Horas declaradas en la distribución que no se programan en la grilla de cursos:
    # docente -> [(cursos, actividad, horas)]
    HORAS_NO_LECTIVAS = {
        'Matemática 2': [('Intv. Media (1° a 4°)', 'Intervención', 24), ('', 'ACLE', 2)],
        'Básica 1': [('', 'Intervención', 7)],
        'Básica 3': [('', 'Intervención', 7)],
        'Básica 4': [('', 'Intervención', 6)],
        'Básica 5': [('', 'Intervención', 6)],
        'Básica 6': [('', 'Intervención', 6)],
        'Básica 7': [('', 'Intervención', 7)],
        'Básica 8': [('', 'Intervención', 6)],
    }

    # Profesores jefes 2026 (filas "Jefatura" de la distribución horaria)
    PROFESORES_JEFES = {
        '1° BÁSICO A': 'Básica 1',
        '1° BÁSICO B': 'Básica 2',
        '2° BÁSICO A': 'Básica 3',
        '2° BÁSICO B': 'Básica 4',
        '3° BÁSICO A': 'Básica 5',
        '3° BÁSICO B': 'Básica 6',
        '4° BÁSICO A': 'Básica 7',
        '4° BÁSICO B': 'Básica 8',
        '5° BÁSICO A': 'Historia 2',
        '5° BÁSICO B': 'Inglés 1',
        '6° BÁSICO A': 'Inglés 4',
        '6° BÁSICO B': 'C. Naturales',
        '7° BÁSICO A': 'C. Naturales y Religión',
        '7° BÁSICO B': 'Inglés 3',
        '8° BÁSICO A': 'Historia 1',
        '8° BÁSICO B': 'Matemática 2',
        '1° MEDIO A': 'Química',
        '1° MEDIO B': 'Inglés 5',
        '2° MEDIO A': 'Matemática 1',
        '2° MEDIO B': 'Lenguaje 4',
        '3° MEDIO A': 'Matemática 3',
        '3° MEDIO B': 'Lenguaje 1',
        '4° MEDIO A': 'Lenguaje 3',
        '4° MEDIO B': 'Inglés 2',  # Supuesto: el documento repite la jefatura de 4°A y omite la de 4°B
    }

    # Etiqueta que muestra cada asignatura en la grilla de los cursos (igual que el Excel 2026)
    ETIQUETAS = {
        'Lenguaje y Comunicación': 'LENGUAJE',
        'Educación Matemática': 'MATEMÁTICA',
        'Idioma Extranjero Inglés': 'INGLÉS',
        'English Skills': 'INGLÉS',
        'Ciencias Sociales': 'C. SOCIALES',
        'Historia, Geografía y CC.SS.': 'HISTORIA',
        'Ciencias Naturales': 'C. NATURALES',
        'Biología': 'BIOLOGÍA',
        'Química': 'QUÍMICA',
        'Física': 'FÍSICA',
        'Educación Tecnológica': 'TECNOLOGÍA',
        'Artes Visuales': 'ARTES',
        'Educación Musical': 'MÚSICA',
        'Artes Visuales / Música': 'ARTES-MÚSICA',
        'Educación Física y Salud': 'ED. FÍSICA',
        'Orientación / Tecnología': 'ORIEN/TEC.',
        'Orientación': 'ORIENTACIÓN',
        'Religión': 'RELIGIÓN',
        'Filosofía': 'FILOSOFÍA',
        'Educación Ciudadana': 'ED. CIUDADANA',
        'Ciencias para la Ciudadanía': 'CC. CIUDADANA',
    }

    # Asignaturas que comparten la regla de "una sesión diaria" (English skills es parte de Inglés)
    FAMILIAS = {'English Skills': 'Idioma Extranjero Inglés'}

    # Asignaturas troncales que se dictan en bloques consecutivos de 90 minutos
    ASIGNATURAS_TRONCALES = ['Matemática', 'Lenguaje', 'Ciencias Naturales', 'Biología', 'Química', 'Física']

    # Salas con capacidad limitada (un curso por bloque)
    RECURSOS_CAPACIDAD = {'Sala English skills': 1}

    # Recintos deportivos: cada uno admite un curso por bloque. El Patio Santo Domingo
    # recibe el desborde de cualquiera de los dos gimnasios (uso penalizado).
    RECINTOS_DEPORTIVOS = ['GIMNASIO A', 'GIMNASIO B', 'PATIO SANTO DOMINGO']

    # Franjas institucionales de electivos (formación diferenciada), fijas según la grilla 2026.
    # Cada franja dura 6 horas: 3 periodos dobles en los que 3° (o 4°) A y B se mezclan.
    # 'periodos': bloque inicial de cada periodo doble; 'etiquetas': texto del 1er y 2do bloque
    # en la grilla del curso; 'grupos': (asignatura, sección, docente, sala, etiqueta en la sala).
    FRANJAS_ELECTIVOS = {
        '3° MEDIO': [
            {
                'periodos': [('Lunes', 3), ('Martes', 7), ('Jueves', 5)],
                'etiquetas': ('CHP 3AB', 'LE-QU-EC S. 1'),
                'grupos': [
                    ('COMPRENSIÓN HISTÓRICA DEL PRESENTE', '3AB', 'Historia 3', 'SALA DE 3°A', 'COMP.H 3AB'),
                    ('LECTURA Y ESCRITURA ESPECIALIZADA', 'S.1', 'Lenguaje 1', 'SALA DE 3°B', 'LECT ESP. S.1'),
                    ('QUÍMICA FORMACIÓN DIFERENCIADA', 'S.1', 'Química', 'ELECTIVO 2', 'QUÍMICA S.1'),
                    ('ECONOMÍA Y SOCIEDAD', 'S.1', 'Historia 1', 'ELECTIVO 3', 'ECONOMÍA S.1'),
                ],
            },
            {
                'periodos': [('Lunes', 7), ('Miércoles', 1), ('Jueves', 3)],
                'etiquetas': ('PROB. EST. 3AB', 'LE-QU-EC S. 2'),
                'grupos': [
                    ('PROBABILIDADES Y ESTADÍSTICA DESCRIPTIVA', '3AB', 'Matemática 1', 'SALA DE 3°A', 'PROB. ES. 3AB'),
                    ('LECTURA Y ESCRITURA ESPECIALIZADA', 'S.2', 'Lenguaje 1', 'SALA DE 3°B', 'LECT ESP. S.2'),
                    ('QUÍMICA FORMACIÓN DIFERENCIADA', 'S.2', 'Química', 'ELECTIVO 2', 'QUÍMICA S.2'),
                    ('ECONOMÍA Y SOCIEDAD', 'S.2', 'Historia 1', 'ELECTIVO 3', 'ECONOMÍA S.2'),
                ],
            },
            {
                'periodos': [('Martes', 5), ('Miércoles', 3), ('Jueves', 7)],
                'etiquetas': ('BIO. CEL. 3AB', 'LE-QU-EC S. 3'),
                'grupos': [
                    ('BIOLOGÍA CELULAR Y MOLECULAR', '3AB', 'Biología', 'SALA DE 3°A', 'BIO. CEL. 3AB'),
                    # El documento indica "3A-B s. 2" para Lenguaje 4; se interpreta como la sección 3
                    ('LECTURA Y ESCRITURA ESPECIALIZADA', 'S.3', 'Lenguaje 4', 'SALA DE 3°B', 'LECT ESP. S.3'),
                    ('QUÍMICA FORMACIÓN DIFERENCIADA', 'S.3', 'Química', 'ELECTIVO 2', 'QUÍMICA S.3'),
                    ('ECONOMÍA Y SOCIEDAD', 'S.3', 'Historia 2', 'ELECTIVO 3', 'ECONOMÍA S.3'),
                ],
            },
        ],
        '4° MEDIO': [
            {
                'periodos': [('Martes', 1), ('Miércoles', 7), ('Viernes', 1)],
                'etiquetas': {'4° MEDIO A': ('DISEÑO 4A', 'LIM-PA-GEO S.1'),
                              '4° MEDIO B': ('DISEÑO 4B', 'LIM-PA-GEO S.1')},
                'grupos': [
                    ('DISEÑO Y ARQUITECTURA', '4A', 'Artes y Tecnología 1', 'SALA DE 4°A', 'DISEÑO 4A'),
                    ('DISEÑO Y ARQUITECTURA', '4B', 'Artes y Tecnología 2', 'SALA DE 4°B', 'DISEÑO 4B'),
                    ('LÍMITES, DERIVADAS E INTEGRALES', 'S.1', 'Matemática 3', 'ELECTIVO 2', 'LIMITES S.1'),
                    ('PARTICIPACIÓN Y ARGUMENTACIÓN EN DEMOCRACIA', 'S.1', 'Lenguaje 1', 'SALA DE TECNOLOGÍA', 'PART.ARG S.1'),
                    ('GEOGRAFÍA, TERRITORIO Y PROBLEMAS SOCIOAMBIENTALES', 'S.1', 'Historia 3', 'ELECTIVO 3', 'GEOGRAFÍA S.1'),
                ],
            },
            {
                'periodos': [('Martes', 3), ('Miércoles', 5), ('Viernes', 3)],
                'etiquetas': {'4° MEDIO A': ('BIO. ECOS 4A', 'LIM-PA-GEO S.2'),
                              '4° MEDIO B': ('BIO. ECOS 4B', 'LIM-PA-GEO S.2')},
                'grupos': [
                    ('BIOLOGÍA DE LOS ECOSISTEMAS', '4A', 'Biología', 'SALA DE 4°A', 'BIO. ECOS. 4A'),
                    ('BIOLOGÍA DE LOS ECOSISTEMAS', '4B', 'C. Naturales', 'ELECTIVO 3', 'BIO. ECOS. 4B'),
                    ('LÍMITES, DERIVADAS E INTEGRALES', 'S.2', 'Matemática 1', 'ELECTIVO 2', 'LIMITES S.2'),
                    ('PARTICIPACIÓN Y ARGUMENTACIÓN EN DEMOCRACIA', 'S.2', 'Lenguaje 1', 'SALA DE TECNOLOGÍA', 'PART.ARG S.2'),
                    ('GEOGRAFÍA, TERRITORIO Y PROBLEMAS SOCIOAMBIENTALES', 'S.2', 'Historia 3', 'SALA DE 4°B', 'GEOGRAFÍA S.2'),
                ],
            },
            {
                'periodos': [('Lunes', 1), ('Jueves', 1), ('Viernes', 5)],
                'etiquetas': ('C. SALUD 4AB', 'LIM-PA-GEO S.3'),
                'grupos': [
                    ('CIENCIAS PARA LA SALUD', '4AB', 'C. Naturales', 'SALA DE 4°B', 'C.SALUD 4AB'),
                    ('LÍMITES, DERIVADAS E INTEGRALES', 'S.3', 'Matemática 1', 'ELECTIVO 2', 'LIMITES S.3'),
                    ('PARTICIPACIÓN Y ARGUMENTACIÓN EN DEMOCRACIA', 'S.3', 'Lenguaje 2', 'SALA DE 4°A', 'PART.ARG S.3'),
                    ('GEOGRAFÍA, TERRITORIO Y PROBLEMAS SOCIOAMBIENTALES', 'S.3', 'Historia 3', 'ELECTIVO 3', 'GEOGRAFÍA S.3'),
                ],
            },
        ],
    }

    # Electivos ofrecidos por nivel (orden de la tabla de horas del Excel 2026)
    ELECTIVOS_OFERTA = {
        '3° MEDIO': [
            'LECTURA Y ESCRITURA ESPECIALIZADA',
            'COMPRENSIÓN HISTÓRICA DEL PRESENTE',
            'PROBABILIDADES Y ESTADÍSTICA DESCRIPTIVA',
            'QUÍMICA FORMACIÓN DIFERENCIADA',
            'BIOLOGÍA CELULAR Y MOLECULAR',
            'ECONOMÍA Y SOCIEDAD',
        ],
        '4° MEDIO': [
            'CIENCIAS PARA LA SALUD',
            'LÍMITES, DERIVADAS E INTEGRALES',
            'PARTICIPACIÓN Y ARGUMENTACIÓN EN DEMOCRACIA',
            'GEOGRAFÍA, TERRITORIO Y PROBLEMAS SOCIOAMBIENTALES',
            'DISEÑO Y ARQUITECTURA',
            'BIOLOGÍA DE LOS ECOSISTEMAS',
        ],
    }

    # Bloqueos de pastoral por sala ('C' = horario de colación). Fuente: hoja SALAS ELECTIVOS 2026
    BLOQUEOS_PASTORAL = {
        'SALA DE TECNOLOGÍA': [('Martes', 'C'), ('Martes', 9), ('Martes', 10),
                               ('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'ELECTIVO 2': [('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'ELECTIVO 3': [('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'SALA PADRE CUETO': [('Martes', 'C'), ('Martes', 9), ('Martes', 10),
                             ('Jueves', 7), ('Jueves', 8), ('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'SALA MADRE PILAR': [('Lunes', 'C'), ('Lunes', 9), ('Lunes', 10),
                             ('Martes', 1), ('Martes', 2), ('Martes', 'C'), ('Martes', 9), ('Martes', 10),
                             ('Jueves', 1), ('Jueves', 2), ('Jueves', 3), ('Jueves', 4),
                             ('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'SALA DE ARTES': [('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
        'PÁRVULOS': [('Jueves', 'C'), ('Jueves', 9), ('Jueves', 10)],
    }

    _malla_cache = None

    @classmethod
    def todos_los_cursos(cls):
        """Párvulos y cursos regulares, en el orden de las hojas del Excel 2026."""
        return cls.CURSOS_PARVULOS + cls.CURSOS

    @classmethod
    def es_parvulo(cls, curso):
        return curso in cls.CURSOS_PARVULOS

    @classmethod
    def nivel(cls, curso):
        """'3° MEDIO A' -> '3° MEDIO'."""
        return curso[:-2]

    @classmethod
    def get_bloques_permitidos(cls, curso, dia):
        """Retorna los bloques pedagógicos de la jornada de un curso en un día."""
        return list(range(1, cls.JORNADAS[curso][cls.DIAS.index(dia)] + 1))

    @classmethod
    def get_total_bloques_semanal(cls, curso):
        """Calcula el total de bloques disponibles en la semana para un curso."""
        return sum(cls.JORNADAS[curso])

    @classmethod
    def zona_deportiva(cls, curso):
        """Gimnasio preferente: B para párvulos y 1° a 4° básico; A para 5° básico a 4° medio."""
        if cls.es_parvulo(curso) or ('BÁSICO' in curso and curso[0] in '1234'):
            return 'GIMNASIO B'
        return 'GIMNASIO A'

    @classmethod
    def codigo_corto(cls, curso):
        """Código del colegio en gimnasios: '1a' (1° básico A), '5A' (5° básico A), '1A' (1° medio A), 'Pk'."""
        if cls.es_parvulo(curso):
            return {'PREKINDER A': 'Pk', 'KINDER A': 'Ka', 'KINDER B': 'Kb'}[curso]
        numero, seccion = curso[0], curso[-1]
        if 'BÁSICO' in curso and numero in '1234':
            return f"{numero}{seccion.lower()}"
        return f"{numero}{seccion}"

    @classmethod
    def codigo_sala(cls, curso):
        """Código del colegio en las salas de especialidad: '6°b' (básica) o '1°B' (media)."""
        numero, seccion = curso[0], curso[-1]
        return f"{numero}°{seccion.lower()}" if 'BÁSICO' in curso else f"{numero}°{seccion}"

    @classmethod
    def familia(cls, asignatura):
        return cls.FAMILIAS.get(asignatura, asignatura)

    @staticmethod
    def docentes_de(docente):
        """Normaliza el docente de una lección a una tupla (las lecciones con co-docencia tienen varios)."""
        return tuple(docente) if isinstance(docente, (tuple, list)) else (docente,)

    @classmethod
    def franjas_de(cls, curso):
        return cls.FRANJAS_ELECTIVOS.get(cls.nivel(curso), [])

    @classmethod
    def etiquetas_franja(cls, franja, curso):
        etiquetas = franja['etiquetas']
        return etiquetas[curso] if isinstance(etiquetas, dict) else etiquetas

    @classmethod
    def nombre_franja(cls, k):
        return f'Formación Diferenciada · Franja {k}'

    @classmethod
    def get_malla_curricular(cls):
        """
        Retorna la asignación curricular completa de cada curso:
        (Asignatura, Horas semanales, Docente(s) asignado(s), Espacio)
        Horas según HORARIO CURSOS 2026.xlsx y docentes según DISTRIBUCIÓN HORARIA 2026.docx.
        Las lecciones con co-docencia (English skills, Artes/Música en media y electivos)
        indican una tupla de docentes que deben estar libres en el mismo bloque.
        """
        if cls._malla_cache is not None:
            return cls._malla_cache

        SK = 'Sala English skills'
        AM = 'Sala de Artes y Música'
        NR = 'C. Naturales y Religión'
        malla = {}

        # ---------------------------------------------------------------------
        # PÁRVULOS (solo Educación Física: 2 bloques simples en días distintos)
        # ---------------------------------------------------------------------
        for c in cls.CURSOS_PARVULOS:
            malla[c] = [('Educación Física y Salud', 2, 'Ed. Física 2', 'Gimnasio')]

        # ---------------------------------------------------------------------
        # 1° BÁSICO A (38 horas)
        # ---------------------------------------------------------------------
        malla['1° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 2', 'Aula'),
            ('Educación Matemática', 8, 'Básica 1', 'Aula'),
            ('Idioma Extranjero Inglés', 6, 'Inglés 1', 'Aula'),
            ('Ciencias Sociales', 3, 'Básica 1', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 1', 'Aula'),
            ('Artes Visuales', 2, 'Básica 1', 'Aula'),
            ('Educación Musical', 2, 'Básica 1', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 1', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 2', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 1° BÁSICO B (38 horas)
        # ---------------------------------------------------------------------
        malla['1° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 2', 'Aula'),
            ('Educación Matemática', 8, 'Básica 1', 'Aula'),
            ('Idioma Extranjero Inglés', 6, 'Inglés 1', 'Aula'),
            ('Ciencias Sociales', 3, 'Básica 2', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 2', 'Aula'),
            ('Artes Visuales', 2, 'Básica 2', 'Aula'),
            ('Educación Musical', 2, 'Básica 2', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 2', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 2', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 2° BÁSICO A (38 horas)
        # ---------------------------------------------------------------------
        malla['2° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 3', 'Aula'),
            ('Educación Matemática', 8, 'Básica 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 1', 'Aula'),
            ('English Skills', 2, ('Inglés 1', 'Inglés 4'), SK),
            ('Ciencias Sociales', 3, 'Básica 3', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 3', 'Aula'),
            ('Artes Visuales', 2, 'Básica 3', 'Aula'),
            ('Educación Musical', 2, 'Básica 3', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 3', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 2', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 2° BÁSICO B (38 horas)
        # ---------------------------------------------------------------------
        malla['2° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 3', 'Aula'),
            ('Educación Matemática', 8, 'Básica 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 1', 'Aula'),
            ('English Skills', 2, ('Inglés 1', 'Inglés 4'), SK),
            ('Ciencias Sociales', 3, 'Básica 4', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 4', 'Aula'),
            ('Artes Visuales', 2, 'Básica 4', 'Aula'),
            ('Educación Musical', 2, 'Básica 4', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 4', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 2', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 3° BÁSICO A (38 horas)
        # ---------------------------------------------------------------------
        malla['3° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 5', 'Aula'),
            ('Educación Matemática', 8, 'Básica 6', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 1', 'Aula'),
            ('English Skills', 2, ('Inglés 1', 'Inglés 3'), SK),
            ('Ciencias Sociales', 3, 'Básica 5', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 5', 'Aula'),
            ('Artes Visuales', 2, 'Básica 5', 'Aula'),
            ('Educación Musical', 2, 'Básica 5', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 5', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 2', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 3° BÁSICO B (38 horas)
        # ---------------------------------------------------------------------
        malla['3° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 5', 'Aula'),
            ('Educación Matemática', 8, 'Básica 6', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('English Skills', 2, ('Inglés 2', 'Inglés 5'), SK),
            ('Ciencias Sociales', 3, 'Básica 6', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 6', 'Aula'),
            ('Artes Visuales', 2, 'Básica 6', 'Aula'),
            ('Educación Musical', 2, 'Básica 6', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 6', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 1', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 4° BÁSICO A (38 horas)
        # ---------------------------------------------------------------------
        malla['4° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 7', 'Aula'),
            ('Educación Matemática', 8, 'Básica 8', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('English Skills', 2, ('Inglés 2', 'Inglés 3'), SK),
            ('Ciencias Sociales', 3, 'Básica 7', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 7', 'Aula'),
            ('Artes Visuales', 2, 'Básica 7', 'Aula'),
            ('Educación Musical', 2, 'Básica 7', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 7', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 1', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 4° BÁSICO B (38 horas)
        # ---------------------------------------------------------------------
        malla['4° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 8, 'Básica 7', 'Aula'),
            ('Educación Matemática', 8, 'Básica 8', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 3', 'Aula'),
            ('English Skills', 2, ('Inglés 3', 'Inglés 5'), SK),
            ('Ciencias Sociales', 3, 'Básica 8', 'Aula'),
            ('Ciencias Naturales', 3, 'Básica 8', 'Aula'),
            ('Artes Visuales', 2, 'Básica 8', 'Aula'),
            ('Educación Musical', 2, 'Básica 8', 'Aula'),
            ('Orientación / Tecnología', 1, 'Básica 8', 'Aula'),
            ('Educación Física y Salud', 3, 'Ed. Física 1', 'Gimnasio'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 5° BÁSICO A (38 horas)
        # ---------------------------------------------------------------------
        malla['5° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 2', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 3', 'Aula'),
            ('English Skills', 2, ('Inglés 3', 'Inglés 4'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 2', 'Aula'),
            ('Ciencias Naturales', 4, 'C. Naturales', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 2', 'Gimnasio'),
            ('Orientación', 1, 'Historia 2', 'Aula'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 5° BÁSICO B (38 horas)
        # ---------------------------------------------------------------------
        malla['5° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 1', 'Aula'),
            ('English Skills', 2, ('Inglés 1', 'Inglés 3'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 2', 'Aula'),
            ('Ciencias Naturales', 4, 'Química', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 1', 'Gimnasio'),
            ('Orientación', 1, 'Inglés 1', 'Aula'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 6° BÁSICO A (37 horas)
        # ---------------------------------------------------------------------
        malla['6° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 6, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 4', 'Aula'),
            ('English Skills', 2, ('Inglés 4', 'Inglés 5'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 2', 'Aula'),
            ('Ciencias Naturales', 4, 'C. Naturales', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 1', 'Gimnasio'),
            ('Orientación', 1, 'Inglés 4', 'Aula'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 6° BÁSICO B (37 horas)
        # ---------------------------------------------------------------------
        malla['6° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 6, 'Matemática 2', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 4', 'Aula'),
            ('English Skills', 2, ('Inglés 4', 'Inglés 2'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 2', 'Aula'),
            ('Ciencias Naturales', 4, 'C. Naturales', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 1', 'Gimnasio'),
            ('Orientación', 1, 'C. Naturales', 'Aula'),
            ('Religión', 2, 'Religión', 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 7° BÁSICO A (37 horas)
        # ---------------------------------------------------------------------
        malla['7° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 2', 'Aula'),
            ('Educación Matemática', 6, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 3', 'Aula'),
            ('English Skills', 2, ('Inglés 3', 'Inglés 5'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 3', 'Aula'),
            ('Ciencias Naturales', 4, NR, 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Educación Tecnológica', 1, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 2', 'Gimnasio'),
            ('Orientación', 1, NR, 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 7° BÁSICO B (37 horas)
        # ---------------------------------------------------------------------
        malla['7° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 2', 'Aula'),
            # El documento repite "7A" para Matemática 4; se interpreta la segunda fila como 7°B
            ('Educación Matemática', 6, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 3', 'Aula'),
            ('English Skills', 2, ('Inglés 3', 'Inglés 5'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 3', 'Aula'),
            ('Ciencias Naturales', 4, NR, 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Educación Tecnológica', 1, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 1', 'Gimnasio'),
            ('Orientación', 1, 'Inglés 3', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 8° BÁSICO A (37 horas)
        # ---------------------------------------------------------------------
        malla['8° BÁSICO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 4', 'Aula'),
            ('Educación Matemática', 6, 'Matemática 3', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('English Skills', 2, ('Inglés 2', 'Inglés 5'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            ('Ciencias Naturales', 4, NR, 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Educación Tecnológica', 1, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Historia 1', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 8° BÁSICO B (37 horas)
        # ---------------------------------------------------------------------
        malla['8° BÁSICO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 4', 'Aula'),
            ('Educación Matemática', 6, 'Matemática 2', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('English Skills', 2, ('Inglés 2', 'Inglés 5'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            ('Ciencias Naturales', 4, 'Biología', 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Educación Tecnológica', 1, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Educación Musical', 2, 'Música', 'Sala de Música'),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Matemática 2', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 1° MEDIO A (40 horas)
        # ---------------------------------------------------------------------
        malla['1° MEDIO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 3', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 5', 'Aula'),
            ('English Skills', 2, ('Inglés 5', 'Inglés 4'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            ('Biología', 4, 'Biología', 'Aula'),
            ('Química', 2, 'Química', 'Aula'),
            ('Física', 2, 'Física', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 1', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Química', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 1° MEDIO B (40 horas)
        # ---------------------------------------------------------------------
        malla['1° MEDIO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 3', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 5', 'Aula'),
            ('English Skills', 2, ('Inglés 5', 'Inglés 4'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            # El documento repite "1A" para Biología; se interpreta la segunda fila como 1° medio B
            ('Biología', 4, 'Biología', 'Aula'),
            ('Química', 2, 'Química', 'Aula'),
            ('Física', 2, 'Física', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 2', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 1', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Inglés 5', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 2° MEDIO A (40 horas)
        # ---------------------------------------------------------------------
        malla['2° MEDIO A'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 4', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 1', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 4', 'Aula'),
            ('English Skills', 2, ('Inglés 4', 'Inglés 3'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            ('Biología', 4, 'Biología', 'Aula'),
            ('Química', 2, 'Química', 'Aula'),
            ('Física', 2, 'Física', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 2', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Matemática 1', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 2° MEDIO B (40 horas)
        # ---------------------------------------------------------------------
        malla['2° MEDIO B'] = [
            ('Lenguaje y Comunicación', 6, 'Lenguaje 2', 'Aula'),
            ('Educación Matemática', 7, 'Matemática 1', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 4', 'Aula'),
            ('English Skills', 2, ('Inglés 4', 'Inglés 3'), SK),
            ('Historia, Geografía y CC.SS.', 4, 'Historia 1', 'Aula'),
            ('Biología', 4, 'Biología', 'Aula'),
            ('Química', 2, 'Química', 'Aula'),
            ('Física', 2, 'Física', 'Aula'),
            ('Educación Tecnológica', 2, 'Artes y Tecnología 1', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 2', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Lenguaje 4', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 3° MEDIO A (42 horas: 24 Plan Común + 18 en 3 franjas de electivos)
        # ---------------------------------------------------------------------
        malla['3° MEDIO A'] = [
            ('Lenguaje y Comunicación', 3, 'Lenguaje 1', 'Aula'),
            ('Educación Matemática', 3, 'Matemática 3', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 5', 'Aula'),
            ('Educación Ciudadana', 2, 'Historia 2', 'Aula'),
            ('Ciencias para la Ciudadanía', 2, 'Biología', 'Aula'),
            ('Filosofía', 2, 'Lenguaje 5', 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 1', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Matemática 3', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 3° MEDIO B (42 horas: 24 Plan Común + 18 en 3 franjas de electivos)
        # ---------------------------------------------------------------------
        malla['3° MEDIO B'] = [
            ('Lenguaje y Comunicación', 3, 'Lenguaje 1', 'Aula'),
            ('Educación Matemática', 3, 'Matemática 3', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 5', 'Aula'),
            ('Educación Ciudadana', 2, 'Historia 2', 'Aula'),
            ('Ciencias para la Ciudadanía', 2, 'Biología', 'Aula'),
            ('Filosofía', 2, 'Lenguaje 5', 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 1', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Lenguaje 1', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 4° MEDIO A (42 horas: 24 Plan Común + 18 en 3 franjas de electivos)
        # ---------------------------------------------------------------------
        malla['4° MEDIO A'] = [
            ('Lenguaje y Comunicación', 3, 'Lenguaje 3', 'Aula'),
            ('Educación Matemática', 3, 'Matemática 1', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('Educación Ciudadana', 2, 'Historia 3', 'Aula'),
            ('Ciencias para la Ciudadanía', 2, 'Química', 'Aula'),
            ('Filosofía', 2, 'Lenguaje 5', 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 1', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Lenguaje 3', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # ---------------------------------------------------------------------
        # 4° MEDIO B (42 horas: 24 Plan Común + 18 en 3 franjas de electivos)
        # ---------------------------------------------------------------------
        malla['4° MEDIO B'] = [
            ('Lenguaje y Comunicación', 3, 'Lenguaje 1', 'Aula'),
            ('Educación Matemática', 3, 'Matemática 4', 'Aula'),
            ('Idioma Extranjero Inglés', 4, 'Inglés 2', 'Aula'),
            ('Educación Ciudadana', 2, 'Historia 3', 'Aula'),
            ('Ciencias para la Ciudadanía', 2, 'Química', 'Aula'),
            ('Filosofía', 2, 'Lenguaje 5', 'Aula'),
            ('Física', 1, 'Física', 'Aula'),
            ('Artes Visuales / Música', 2, ('Artes y Tecnología 2', 'Música'), AM),
            ('Educación Física y Salud', 2, 'Ed. Física 3', 'Gimnasio'),
            ('Orientación', 1, 'Inglés 2', 'Aula'),
            ('Religión', 2, NR, 'Aula'),
        ]

        # Franjas de electivos de 3° y 4° medio (6 horas cada una, docentes de todas sus secciones)
        for nivel, franjas in cls.FRANJAS_ELECTIVOS.items():
            for k, franja in enumerate(franjas, 1):
                docentes = tuple(dict.fromkeys(g[2] for g in franja['grupos']))
                for seccion in ('A', 'B'):
                    malla[f'{nivel} {seccion}'].append((cls.nombre_franja(k), 6, docentes, 'Salas electivos'))

        cls._malla_cache = malla
        return malla

    @classmethod
    def carga_docente(cls):
        """Horas de grilla por docente según la malla (los electivos se cuentan una vez por franja)."""
        carga = defaultdict(int)
        for curso, lecciones in cls.get_malla_curricular().items():
            for asig, horas, doc, _ in lecciones:
                if asig.startswith('Formación Diferenciada'):
                    continue
                for d in cls.docentes_de(doc):
                    carga[d] += horas
        for franjas in cls.FRANJAS_ELECTIVOS.values():
            for franja in franjas:
                for d in dict.fromkeys(g[2] for g in franja['grupos']):
                    carga[d] += 6
        return carga


# =============================================================================
# 2. ESTRUCTURA DE DATOS DEL HORARIO ESCOLAR
# =============================================================================

class HorarioEscolar:
    """Representación matricial global del horario del colegio."""

    def __init__(self):
        # asignaciones[curso][dia][bloque] = {asignatura, docentes, docente, espacio, etiqueta}
        self.asignaciones = {
            c: {d: {} for d in DatosColegio.DIAS} for c in DatosColegio.todos_los_cursos()
        }
        # docente_ocupado[docente][(dia, bloque)] = (curso, asignatura)
        self.docente_ocupado = defaultdict(dict)
        # docente_claves[docente][(dia, bloque)] = {cursos o grupos atendidos en ese bloque}
        self.docente_claves = defaultdict(lambda: defaultdict(set))
        # docente_detalle[docente][(dia, bloque)] = {'etiqueta', 'sala'} (electivos)
        self.docente_detalle = defaultdict(dict)
        # gimnasio_ocupado[(dia, bloque)] = list(cursos)
        self.gimnasio_ocupado = defaultdict(list)

    def esta_disponible(self, curso, dia, bloque, docente, es_gimnasio=False):
        """Evalúa si un bloque está libre para el curso y el docente, y si hay recinto deportivo."""
        if bloque not in DatosColegio.get_bloques_permitidos(curso, dia):
            return False
        if bloque in self.asignaciones[curso][dia]:
            return False
        if (dia, bloque) in self.docente_ocupado[docente]:
            return False
        if es_gimnasio and len(self.gimnasio_ocupado[(dia, bloque)]) >= len(DatosColegio.RECINTOS_DEPORTIVOS):
            return False
        return True

    def asignar(self, curso, dia, bloque, asignatura, docentes, espacio='Aula', etiqueta=None,
                clave=None, detalle=None):
        """
        Asigna una lección en el horario garantizando actualización de estados.
        'clave' identifica a quién atiende el docente (el curso, o el grupo compartido de un
        electivo) y 'detalle' permite indicar por docente la etiqueta y sala de su sección.
        """
        docentes = DatosColegio.docentes_de(docentes)
        self.asignaciones[curso][dia][bloque] = {
            'asignatura': asignatura,
            'docentes': docentes,
            'docente': ' / '.join(docentes),
            'espacio': espacio,
            'etiqueta': etiqueta or DatosColegio.ETIQUETAS.get(asignatura, asignatura.upper()),
        }
        clave = clave or curso
        for doc in docentes:
            self.docente_claves[doc][(dia, bloque)].add(clave)
            if detalle and doc in detalle:
                etiqueta_doc, sala, asig_doc = detalle[doc]
                self.docente_ocupado[doc][(dia, bloque)] = (clave, asig_doc)
                self.docente_detalle[doc][(dia, bloque)] = {'etiqueta': etiqueta_doc, 'sala': sala}
            else:
                self.docente_ocupado[doc][(dia, bloque)] = (curso, asignatura)
        if espacio == 'Gimnasio':
            self.gimnasio_ocupado[(dia, bloque)].append(curso)

    def desasignar(self, curso, dia, bloque):
        """Elimina una asignación."""
        if bloque in self.asignaciones[curso][dia]:
            item = self.asignaciones[curso][dia].pop(bloque)
            for doc in item['docentes']:
                self.docente_claves[doc][(dia, bloque)].discard(curso)
                if not self.docente_claves[doc][(dia, bloque)]:
                    self.docente_ocupado[doc].pop((dia, bloque), None)
                    self.docente_detalle[doc].pop((dia, bloque), None)
            if item.get('espacio') == 'Gimnasio':
                if curso in self.gimnasio_ocupado[(dia, bloque)]:
                    self.gimnasio_ocupado[(dia, bloque)].remove(curso)

    def asignar_recintos(self):
        """
        Distribuye los cursos en Ed. Física de cada bloque entre los recintos deportivos:
        el primero de cada zona usa su gimnasio y el desborde va al Patio Santo Domingo.
        Retorna {(recinto, dia, bloque): [cursos]}.
        """
        recintos = defaultdict(list)
        for (dia, bloque), cursos in self.gimnasio_ocupado.items():
            for curso in cursos:
                zona = DatosColegio.zona_deportiva(curso)
                destino = zona if not recintos[(zona, dia, bloque)] else 'PATIO SANTO DOMINGO'
                recintos[(destino, dia, bloque)].append(curso)
        return recintos


# =============================================================================
# 3. VALIDADOR DE RESTRICCIONES (AUDITORÍA 100% EXHAUSTIVA)
# =============================================================================

class ValidadorRestricciones:
    """Verifica el estricto cumplimiento de las restricciones duras y los criterios de calidad."""

    @staticmethod
    def auditar(horario):
        malla = DatosColegio.get_malla_curricular()
        reporte = {
            'valido': True,
            'errores_duros': [],
            'advertencias': [],
            'metricas': {}
        }
        errores = reporte['errores_duros']
        conteo = defaultdict(int)

        # HC 1, 2: Choques docentes (Clash-Free Teacher). Un electivo compartido por A y B
        # cuenta como un solo grupo.
        colisiones_docentes = 0
        for doc, bloques in horario.docente_claves.items():
            for (dia, bloque), claves in bloques.items():
                if len(claves) > 1:
                    colisiones_docentes += 1
                    errores.append(f"Colisión Docente: {doc} atiende {sorted(claves)} el {dia} bloque {bloque}")
        conteo['colisiones_docentes'] = colisiones_docentes

        # HC 3: Ocupación única de cada curso dentro de su jornada
        for curso in DatosColegio.todos_los_cursos():
            for dia in DatosColegio.DIAS:
                permitidos = DatosColegio.get_bloques_permitidos(curso, dia)
                for bloque in horario.asignaciones[curso][dia]:
                    if bloque not in permitidos:
                        errores.append(f"Bloque fuera de jornada: {curso} el {dia} bloque {bloque}")
                        conteo['fuera_jornada'] += 1

        # HC 4: Cumplimiento estricto de cargas horarias (y jornada completa sin huecos)
        horas_por_materia = defaultdict(lambda: defaultdict(int))
        total_asignados = defaultdict(int)
        for curso in DatosColegio.todos_los_cursos():
            for dia in DatosColegio.DIAS:
                for bloque, item in horario.asignaciones[curso][dia].items():
                    horas_por_materia[curso][item['asignatura']] += 1
                    total_asignados[curso] += 1
        esperado_global = 0
        for curso, materias in malla.items():
            esperado = sum(h for _, h, _, _ in materias)
            esperado_global += esperado
            if total_asignados[curso] != esperado:
                errores.append(f"Carga horaria incorrecta en {curso}: requerida={esperado}, asignada={total_asignados[curso]}")
                conteo['cargas'] += 1
            if not DatosColegio.es_parvulo(curso) and esperado != DatosColegio.get_total_bloques_semanal(curso):
                errores.append(f"La malla de {curso} ({esperado} h) no coincide con su jornada "
                               f"({DatosColegio.get_total_bloques_semanal(curso)} bloques)")
                conteo['cargas'] += 1
            for asig, hrs, _, _ in materias:
                if horas_por_materia[curso][asig] != hrs:
                    errores.append(f"{curso} - {asig}: horas esperadas {hrs} != asignadas {horas_por_materia[curso][asig]}")
                    conteo['cargas'] += 1

        # HC 5: Orientación y jefatura a cargo del profesor jefe de cada curso
        for curso in DatosColegio.CURSOS:
            jefe = DatosColegio.PROFESORES_JEFES[curso]
            for dia in DatosColegio.DIAS:
                for bloque, item in horario.asignaciones[curso][dia].items():
                    if item['asignatura'] in ('Orientación', 'Orientación / Tecnología') and jefe not in item['docentes']:
                        errores.append(f"Orientación/Jefatura en {curso} no está a cargo de {jefe}")
                        conteo['jefaturas'] += 1

        # HC 6: Sincronización de las franjas de electivos (3° y 4° medio A y B)
        for nivel in DatosColegio.FRANJAS_ELECTIVOS:
            ocupacion = []
            for seccion in ('A', 'B'):
                curso = f"{nivel} {seccion}"
                ocupacion.append({
                    (dia, b, item['asignatura'])
                    for dia in DatosColegio.DIAS
                    for b, item in horario.asignaciones[curso][dia].items()
                    if item['asignatura'].startswith('Formación Diferenciada')
                })
            if ocupacion[0] != ocupacion[1]:
                errores.append(f"Desincronización de electivos en {nivel}: {sorted(ocupacion[0] ^ ocupacion[1])}")
                conteo['electivos'] += 1

        # HC 7: Recintos deportivos (Gimnasio A, Gimnasio B y Patio Santo Domingo, un curso cada uno)
        uso_patio = 0
        for (recinto, dia, bloque), cursos in horario.asignar_recintos().items():
            if recinto == 'PATIO SANTO DOMINGO':
                uso_patio += 1
            if len(cursos) > 1:
                errores.append(f"Capacidad de {recinto} excedida el {dia} bloque {bloque}: {cursos}")
                conteo['recintos'] += 1

        # HC 8: Salas de capacidad limitada (English skills) y co-docencia completa
        uso_salas = defaultdict(list)
        for curso in DatosColegio.CURSOS:
            for dia in DatosColegio.DIAS:
                for bloque, item in horario.asignaciones[curso][dia].items():
                    if item['espacio'] in DatosColegio.RECURSOS_CAPACIDAD:
                        uso_salas[(item['espacio'], dia, bloque)].append(curso)
                    if item['asignatura'] == 'English Skills' and len(item['docentes']) < 2:
                        errores.append(f"English skills sin co-docencia en {curso} el {dia} bloque {bloque}")
                        conteo['salas'] += 1
        for (sala, dia, bloque), cursos in uso_salas.items():
            if len(cursos) > DatosColegio.RECURSOS_CAPACIDAD[sala]:
                errores.append(f"{sala} ocupada por {cursos} el {dia} bloque {bloque}")
                conteo['salas'] += 1

        # HC 9: Una sola sesión diaria por asignatura (una sesión = 1 bloque o un par consecutivo)
        # HC 10: Las troncales nunca quedan en bloques separados del mismo día
        same_day_violations = 0
        core_non_consecutive = 0
        for curso in DatosColegio.todos_los_cursos():
            por_dia = defaultdict(list)
            for dia in DatosColegio.DIAS:
                for bloque, item in horario.asignaciones[curso][dia].items():
                    if not item['asignatura'].startswith('Formación Diferenciada'):
                        por_dia[(dia, DatosColegio.familia(item['asignatura']))].append(bloque)
            for (dia, fam), bloques in por_dia.items():
                bl = sorted(bloques)
                es_par = len(bl) == 2 and (bl[0], bl[1]) in DatosColegio.PARES
                if len(bl) > 2 or (len(bl) == 2 and not es_par):
                    same_day_violations += 1
                    errores.append(f"Asignatura repetida en {curso} el {dia}: '{fam}' en bloques {bl}")
                    if any(t.lower() in fam.lower() for t in DatosColegio.ASIGNATURAS_TRONCALES):
                        core_non_consecutive += 1

        # HC 11: Las salas usadas no coinciden con los bloqueos de pastoral
        for nivel, franjas in DatosColegio.FRANJAS_ELECTIVOS.items():
            for franja in franjas:
                for dia, b0 in franja['periodos']:
                    for grupo in franja['grupos']:
                        bloqueos = DatosColegio.BLOQUEOS_PASTORAL.get(grupo[3], [])
                        if (dia, b0) in bloqueos or (dia, b0 + 1) in bloqueos:
                            errores.append(f"{grupo[3]} bloqueada por pastoral el {dia} ({grupo[4]})")
                            conteo['pastoral'] += 1

        # MÉTRICAS BLANDAS: Ventanas docentes y bloques dobles
        total_ventanas = 0
        for doc, bloques in horario.docente_ocupado.items():
            for dia in DatosColegio.DIAS:
                ocupados = sorted(b for (d, b) in bloques if d == dia)
                for i in range(len(ocupados) - 1):
                    total_ventanas += ocupados[i + 1] - ocupados[i] - 1

        total_sesiones = 0
        bloques_dobles_logrados = 0
        for curso in DatosColegio.CURSOS:
            for dia in DatosColegio.DIAS:
                asig = horario.asignaciones[curso][dia]
                bloques = sorted(asig)
                i = 0
                while i < len(bloques):
                    b = bloques[i]
                    if i + 1 < len(bloques) and bloques[i + 1] == b + 1 and asig[b]['asignatura'] == asig[b + 1]['asignatura']:
                        bloques_dobles_logrados += 2
                        total_sesiones += 2
                        i += 2
                        continue
                    total_sesiones += 1
                    i += 1

        pct_dobles = (bloques_dobles_logrados / total_sesiones * 100) if total_sesiones > 0 else 0
        if uso_patio:
            reporte['advertencias'].append(f"Se usa el Patio Santo Domingo en {uso_patio} bloques de Ed. Física")

        reporte['valido'] = len(errores) == 0
        reporte['metricas'] = {
            'total_bloques_asignados': sum(total_asignados.values()),
            'total_bloques_esperados': esperado_global,
            'total_ventanas_docentes': total_ventanas,
            'porcentaje_bloques_dobles': round(pct_dobles, 1),
            'cursos_auditados': len(DatosColegio.todos_los_cursos()),
            'docentes_auditados': len(horario.docente_ocupado),
            'colisiones_docentes': conteo['colisiones_docentes'],
            'errores_carga': conteo['cargas'] + conteo['fuera_jornada'],
            'errores_jefatura': conteo['jefaturas'],
            'errores_electivos': conteo['electivos'],
            'errores_recintos': conteo['recintos'],
            'errores_salas': conteo['salas'] + conteo['pastoral'],
            'uso_patio': uso_patio,
            'mismo_dia_violaciones': same_day_violations,
            'core_no_consecutivo': core_non_consecutive
        }

        return reporte


# =============================================================================
# 4. MOTOR ALGORÍTMICO DE ASIGNACIÓN (HEURÍSTICA CONSTRUCTIVA + MIN-CONFLICTS TABÚ)
# =============================================================================

class MotorHorarios:
    """
    Motor heurístico para resolver el School Timetabling Problem del Colegio MMDD.

    Cada curso se descompone en sesiones (bloques dobles de 90 min o simples de 45 min).
    Los dobles ocupan siempre un par pedagógico (1-2, 3-4, 5-6, 7-8, 9-10), por lo que
    nunca cruzan un recreo, y la jornada de cada curso queda completa sin huecos.
    Las franjas de electivos son fijas y compartidas por las secciones A y B.

    1. Construcción voraz: sesiones dobles y luego simples, de la más restringida a la
       menos restringida, en la posición de menor costo.
    2. Búsqueda local Min-Conflicts: se elige una restricción violada, una de sus sesiones
       y el mejor intercambio (par por par, o bloque simple por bloque simple) dentro del
       curso, con lista tabú, ruido aleatorio y reinicios desde la mejor solución.
    """

    PESO_DURO = 1000   # Costo de cada violación dura
    PESO_PATIO = 1     # Costo blando por bloque de Ed. Física desbordado al Patio Santo Domingo

    def __init__(self, seed=42, max_pasos=150000):
        self.seed = seed
        self.max_pasos = max_pasos

    def generar(self):
        """
        Ejecuta la construcción y optimización de horarios.
        Busca 0 colisiones duras, 0 repeticiones de asignaturas por día y bloques dobles.
        """
        rng = random.Random(self.seed)
        self._preparar()
        self._construir(rng)
        self._buscar(rng)
        return self._a_horario()

    # ------------------------------------------------------------------
    # Modelo de sesiones
    # ------------------------------------------------------------------
    def _preparar(self):
        malla = DatosColegio.get_malla_curricular()
        self.cursos = DatosColegio.todos_los_cursos()
        self.ses = []                                 # sesiones (movibles y fijas)
        self.pos = {}                                 # sid -> tupla de celdas (dia, bloque)
        self.celda = {c: {} for c in self.cursos}     # celdas movibles: (dia, bloque) -> sid | None
        self.pares = {c: [] for c in self.cursos}     # pares pedagógicos completos y movibles
        self.par_de = {c: {} for c in self.cursos}    # celda -> par que la contiene
        self.occ = defaultdict(set)                   # clave de restricción -> sids
        self.malos = set()                            # claves con violación dura
        self.blandos = set()                          # claves con costo blando
        self.costo = 0
        self.duras = 0

        # Franjas de electivos: sesiones fijas compartidas por A y B
        fijas = defaultdict(set)
        for nivel, franjas in DatosColegio.FRANJAS_ELECTIVOS.items():
            for k, franja in enumerate(franjas, 1):
                docentes = tuple(dict.fromkeys(g[2] for g in franja['grupos']))
                for dia, b0 in franja['periodos']:
                    celdas = ((dia, b0), (dia, b0 + 1))
                    sid = self._nueva_sesion(None, DatosColegio.nombre_franja(k), docentes, 'Salas electivos',
                                             2, fijo=True, clave=f'ELECTIVOS {nivel}', nivel=nivel, franja=k)
                    self.pos[sid] = celdas
                    for seccion in ('A', 'B'):
                        fijas[f'{nivel} {seccion}'].update(celdas)

        # Celdas y pares movibles de cada curso
        for c in self.cursos:
            for dia in DatosColegio.DIAS:
                permitidos = DatosColegio.get_bloques_permitidos(c, dia)
                for b in permitidos:
                    if (dia, b) not in fijas[c]:
                        self.celda[c][(dia, b)] = None
                for b1, b2 in DatosColegio.PARES:
                    if (dia, b1) in self.celda[c] and (dia, b2) in self.celda[c]:
                        par = ((dia, b1), (dia, b2))
                        self.pares[c].append(par)
                        self.par_de[c][(dia, b1)] = par
                        self.par_de[c][(dia, b2)] = par

        # Sesiones de cada curso: dobles para las horas pares y un simple para el resto
        for c in self.cursos:
            for asig, horas, doc, espacio in malla[c]:
                if asig.startswith('Formación Diferenciada'):
                    continue
                docentes = DatosColegio.docentes_de(doc)
                if DatosColegio.es_parvulo(c):
                    largos = [1] * horas
                else:
                    largos = [2] * (horas // 2) + [1] * (horas % 2)
                for largo in largos:
                    self._nueva_sesion(c, asig, docentes, espacio, largo)

        for sid, s in enumerate(self.ses):
            if s['fijo']:
                self._registrar(sid, self.pos[sid], +1)

    def _nueva_sesion(self, curso, asig, docentes, espacio, largo, fijo=False, clave=None, **extra):
        sesion = {
            'curso': curso,
            'asig': asig,
            'fam': None if fijo else DatosColegio.familia(asig),
            'docs': docentes,
            'espacio': espacio,
            'recurso': espacio if espacio in DatosColegio.RECURSOS_CAPACIDAD else None,
            'zona': DatosColegio.zona_deportiva(curso) if espacio == 'Gimnasio' else None,
            'largo': largo,
            'fijo': fijo,
            'clave': clave or curso,
        }
        sesion.update(extra)
        self.ses.append(sesion)
        return len(self.ses) - 1

    # ------------------------------------------------------------------
    # Costos incrementales
    # ------------------------------------------------------------------
    def _costo_clave(self, k):
        """Retorna (costo, violaciones duras, costo blando) de una clave de restricción."""
        s = self.occ.get(k)
        if not s or len(s) < 2:
            return 0, 0, 0
        tipo = k[0]
        if tipo == 'D':
            n = len({self.ses[i]['clave'] for i in s}) - 1
            return n * self.PESO_DURO, n, 0
        if tipo == 'F':
            n = len(s) - 1
            return n * self.PESO_DURO, n, 0
        if tipo == 'R':
            n = max(0, len(s) - DatosColegio.RECURSOS_CAPACIDAD[k[1]])
            return n * self.PESO_DURO, n, 0
        # tipo 'G': un curso por gimnasio y uno más en el Patio Santo Domingo
        zona_b = sum(1 for i in s if self.ses[i]['zona'] == 'GIMNASIO B')
        exceso = max(0, zona_b - 1) + max(0, len(s) - zona_b - 1)
        dura = max(0, exceso - 1)
        blando = min(exceso, 1) * self.PESO_PATIO
        return dura * self.PESO_DURO + blando, dura, blando

    def _claves_sesion(self, sid, celdas):
        s = self.ses[sid]
        claves = []
        for (dia, b) in celdas:
            for doc in s['docs']:
                claves.append(('D', doc, dia, b))
            if s['zona']:
                claves.append(('G', dia, b))
            if s['recurso']:
                claves.append(('R', s['recurso'], dia, b))
        if s['fam'] is not None:
            claves.append(('F', s['curso'], celdas[0][0], s['fam']))
        return claves

    def _registrar(self, sid, celdas, signo):
        """Agrega (+1) o quita (-1) una sesión de sus celdas y retorna la variación de costo."""
        delta = 0
        for k in self._claves_sesion(sid, celdas):
            c0, d0, b0 = self._costo_clave(k)
            if signo > 0:
                self.occ[k].add(sid)
            else:
                self.occ[k].discard(sid)
            c1, d1, b1 = self._costo_clave(k)
            delta += c1 - c0
            self.duras += d1 - d0
            if d1 > 0:
                self.malos.add(k)
            else:
                self.malos.discard(k)
            if b1 > 0:
                self.blandos.add(k)
            else:
                self.blandos.discard(k)
        self.costo += delta
        return delta

    def _mover(self, c, mapa):
        """Permuta el contenido de las celdas de un curso según 'mapa' y retorna la variación de costo."""
        cel = self.celda[c]
        sids = []
        for x in mapa:
            sid = cel[x]
            if sid is not None and sid not in sids:
                sids.append(sid)
        delta = 0
        nuevas = {}
        for sid in sids:
            delta += self._registrar(sid, self.pos[sid], -1)
            nuevas[sid] = tuple(sorted(mapa[x] for x in self.pos[sid]))
        contenido = {x: cel[x] for x in mapa}
        for x, y in mapa.items():
            cel[y] = contenido[x]
        for sid in sids:
            self.pos[sid] = nuevas[sid]
            delta += self._registrar(sid, nuevas[sid], +1)
        return delta

    # ------------------------------------------------------------------
    # Fase 1: construcción voraz
    # ------------------------------------------------------------------
    def _construir(self, rng):
        carga = DatosColegio.carga_docente()
        movibles = [i for i, s in enumerate(self.ses) if not s['fijo']]
        orden = sorted(movibles, key=lambda i: (
            -self.ses[i]['largo'],
            -len(self.ses[i]['docs']),
            -max(carga[d] for d in self.ses[i]['docs']),
            self.ses[i]['zona'] is None,
            rng.random(),
        ))
        for sid in orden:
            s = self.ses[sid]
            c = s['curso']
            if s['largo'] == 2:
                opciones = [par for par in self.pares[c] if self.celda[c][par[0]] is None and self.celda[c][par[1]] is None]
            else:
                opciones = [(x,) for x, v in self.celda[c].items() if v is None]
            mejor, mejor_delta = None, None
            rng.shuffle(opciones)
            for celdas in opciones:
                delta = self._registrar(sid, celdas, +1)
                self._registrar(sid, celdas, -1)
                if mejor_delta is None or delta < mejor_delta:
                    mejor, mejor_delta = celdas, delta
                    if delta == 0:
                        break
            if mejor is None:
                raise RuntimeError(f"No hay espacio en la jornada de {c} para {s['asig']}")
            for x in mejor:
                self.celda[c][x] = sid
            self.pos[sid] = mejor
            self._registrar(sid, mejor, +1)

    # ------------------------------------------------------------------
    # Fase 2: búsqueda local Min-Conflicts con lista tabú
    # ------------------------------------------------------------------
    def _vecinos(self, sid):
        """Intercambios válidos que mueven la sesión 'sid' dentro de su curso."""
        s = self.ses[sid]
        c = s['curso']
        cel = self.celda[c]
        celdas = self.pos[sid]
        movimientos = []
        propio = self.par_de[c].get(celdas[0])
        if propio is not None:
            for otro in self.pares[c]:
                if otro != propio:
                    movimientos.append({propio[0]: otro[0], propio[1]: otro[1], otro[0]: propio[0], otro[1]: propio[1]})
        if s['largo'] == 1:
            u = celdas[0]
            for v, sv in cel.items():
                if v != u and (sv is None or self.ses[sv]['largo'] == 1):
                    movimientos.append({u: v, v: u})
        return movimientos

    def _buscar(self, rng):
        tabu = {}
        mejor_costo = self.costo
        mejor_pos = dict(self.pos)
        sin_mejora = 0
        for paso in range(self.max_pasos):
            if self.costo == 0:
                break
            if self.duras == 0 and sin_mejora > 3000:
                break   # factible y sin mejoras en el uso del patio
            fuente = self.malos if self.malos else self.blandos
            if not fuente:
                break
            k = rng.choice(sorted(fuente))
            candidatos = sorted(i for i in self.occ[k] if not self.ses[i]['fijo'])
            if not candidatos:
                break
            sid = rng.choice(candidatos)
            c = self.ses[sid]['curso']
            movimientos = self._vecinos(sid)
            if not movimientos:
                continue

            if rng.random() < 0.03:
                elegido = rng.choice(movimientos)
            else:
                elegido, mejor_delta = None, None
                for mov in movimientos:
                    destino = tuple(sorted(mov[x] for x in self.pos[sid]))
                    delta = self._mover(c, mov)
                    self._mover(c, {y: x for x, y in mov.items()})
                    es_tabu = tabu.get((sid, destino), -1) > paso
                    if es_tabu and self.costo + delta >= mejor_costo:
                        continue
                    if mejor_delta is None or delta < mejor_delta or (delta == mejor_delta and rng.random() < 0.5):
                        elegido, mejor_delta = mov, delta
                if elegido is None:
                    continue

            origen = self.pos[sid]
            self._mover(c, elegido)
            tabu[(sid, origen)] = paso + 10 + rng.randint(0, 10)

            if self.costo < mejor_costo:
                mejor_costo = self.costo
                mejor_pos = dict(self.pos)
                sin_mejora = 0
            else:
                sin_mejora += 1
                if sin_mejora % 2500 == 0:
                    # Reinicio: volver a la mejor solución y perturbarla
                    self._restaurar(mejor_pos)
                    for _ in range(15):
                        c2 = rng.choice(self.cursos)
                        sids_c = sorted({v for v in self.celda[c2].values() if v is not None})
                        if sids_c:
                            vecinos = self._vecinos(rng.choice(sids_c))
                            if vecinos:
                                self._mover(c2, rng.choice(vecinos))
        if self.costo > mejor_costo:
            self._restaurar(mejor_pos)

    def _restaurar(self, posiciones):
        """Reconstruye todas las estructuras a partir de una asignación de posiciones."""
        self.occ = defaultdict(set)
        self.malos = set()
        self.blandos = set()
        self.costo = 0
        self.duras = 0
        self.pos = dict(posiciones)
        for c in self.cursos:
            for x in self.celda[c]:
                self.celda[c][x] = None
        for sid, s in enumerate(self.ses):
            if not s['fijo']:
                for x in self.pos[sid]:
                    self.celda[s['curso']][x] = sid
            self._registrar(sid, self.pos[sid], +1)

    # ------------------------------------------------------------------
    # Salida
    # ------------------------------------------------------------------
    def _a_horario(self):
        horario = HorarioEscolar()
        for sid, s in enumerate(self.ses):
            if not s['fijo']:
                for dia, b in self.pos[sid]:
                    horario.asignar(s['curso'], dia, b, s['asig'], s['docs'], s['espacio'])
                continue
            franja = DatosColegio.FRANJAS_ELECTIVOS[s['nivel']][s['franja'] - 1]
            detalle = {}
            for asig, seccion, doc, sala, etiqueta in franja['grupos']:
                detalle[doc] = (etiqueta, sala, f"{asig.capitalize()} ({seccion})")
            for i, (dia, b) in enumerate(self.pos[sid]):
                for seccion in ('A', 'B'):
                    curso = f"{s['nivel']} {seccion}"
                    etiqueta = DatosColegio.etiquetas_franja(franja, curso)[i]
                    horario.asignar(curso, dia, b, s['asig'], s['docs'], s['espacio'], etiqueta=etiqueta,
                                    clave=f"{s['nivel']} A-B", detalle=detalle)
        return horario


# =============================================================================
# 5. EXPORTADOR DE PLANILLAS EXCEL (.XLSX) CON EL DISEÑO OFICIAL 2026
# =============================================================================

class ExportadorExcel:
    """
    Genera los libros .xlsx replicando el diseño de 'data/Data 2026/HORARIO CURSOS 2026.xlsx':
    fuente Cavolini 8, encabezados y recreos en amarillo, número de bloque y horario en crema,
    bordes finos, grilla en las columnas B-H y tabla de horas por asignatura bajo cada grilla.
    """

    FUENTE = 'Cavolini'
    AMARILLO = 'FFFFE598'   # Encabezados, recreos, colación y nombres de asignaturas
    CREMA = 'FFFEF2CB'      # Número de bloque, horario y bloques fuera de la jornada
    BLANCO = 'FFFFFFFF'     # Celdas de clases
    SKILLS = 'FFFFFF00'     # English skills
    PASTORAL = 'FFFF9900'   # Bloqueos de pastoral
    BORDE = 'FFED7D31'      # Bordes finos naranjos de todas las grillas y tablas
    BORDE_MEDIA = 'FF0070C0'  # Bordes azules de las asignaturas en la tabla de 3° y 4° medio
    BORDE_DOCUMENTO = 'FF000000'  # Bordes negros de la distribución horaria (documento Word 2026)
    GRIS = 'FFC0C0C0'       # Encabezado de la distribución horaria (documento Word 2026)
    CELESTE = 'FFB7DEE8'    # Jefaturas en la distribución horaria
    VERDE = 'FFC4D79B'      # English skills en la distribución horaria
    AZUL = 'FFBDD6EE'       # Tecn/C.Orien. en la distribución horaria

    ANCHOS = {'A': 0.66, 'B': 8.66, 'C': 12.66, 'D': 13.89, 'E': 14.44,
              'F': 14.44, 'G': 14.44, 'H': 14.44, 'I': 10.66}

    # Excel oficial del que se toma el escudo del colegio (si no existe, las hojas van sin escudo)
    RUTA_EXCEL_OFICIAL = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                      'data', 'Data 2026', 'HORARIO CURSOS 2026.xlsx')
    _escudo = None

    # Disposición de la hoja SALAS ELECTIVOS: (fila del título, columna inicial, sala, subtítulo)
    DISPOSICION_SALAS = [
        (2, 2, 'ELECTIVO 1', 'INGLÉS - English skills'),
        (2, 10, 'SALA DE TECNOLOGÍA', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (2, 18, 'SALA DE 4°A', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (21, 2, 'ELECTIVO 2', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (21, 10, 'SALA PADRE CUETO', 'SALA PASTORAL'),
        (21, 18, 'SALA DE 4°B', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (40, 2, 'ELECTIVO 3', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (40, 10, 'SALA MADRE PILAR', 'SALA PASTORAL'),
        (59, 2, 'SALA DE ARTES', 'ARTES VISUALES (MEDIA)'),
        (59, 10, 'SALA DE 3°A', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (78, 2, 'SALA DE MÚSICA', 'MÚSICA'),
        (78, 10, 'SALA DE 3°B', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
    ]

    # Nombre corto de cada asignatura en la distribución horaria (como en el documento Word 2026)
    NOMBRES_DISTRIBUCION = {
        'Lenguaje y Comunicación': 'Lenguaje',
        'Educación Matemática': 'Matemática',
        'Idioma Extranjero Inglés': 'Inglés',
        'Ciencias Sociales': 'C. sociales',
        'Historia, Geografía y CC.SS.': 'Historia',
        'Educación Tecnológica': 'Tecnología',
        'Artes Visuales': 'Artes',
        'Educación Musical': 'Música',
        'Educación Física y Salud': 'Ed. Física',
        'Educación Ciudadana': 'Educación ciudadana',
        'Ciencias para la Ciudadanía': 'Ciencias para la ciudadanía',
    }

    # ------------------------------------------------------------------
    # Utilidades de formato
    # ------------------------------------------------------------------
    @classmethod
    def _nuevo_libro(cls):
        """Libro vacío cuya fuente por defecto es Cavolini 11, igual que el Excel 2026."""
        if openpyxl is None:
            raise RuntimeError("La exportación a Excel requiere el paquete openpyxl (pip install openpyxl).")
        wb = openpyxl.Workbook()
        wb._fonts = IndexedList([Font(name=cls.FUENTE, sz=11, family=2)])
        return wb

    @classmethod
    def _insertar_escudo(cls, ws):
        """Inserta el escudo del colegio en B1 (70 x 93 px), tomado del Excel oficial 2026."""
        if cls._escudo is None:
            cls._escudo = b''
            try:
                with zipfile.ZipFile(cls.RUTA_EXCEL_OFICIAL) as z:
                    medios = sorted(n for n in z.namelist() if n.startswith('xl/media/'))
                    if medios:
                        cls._escudo = z.read(medios[0])
            except (OSError, zipfile.BadZipFile):
                pass
        if not cls._escudo:
            return
        try:
            imagen = ImagenExcel(io.BytesIO(cls._escudo))
        except ImportError:   # openpyxl requiere Pillow para insertar imágenes; se omite el escudo
            return
        imagen.width, imagen.height = 70, 93
        ws.add_image(imagen, 'B1')

    @classmethod
    def _estilo(cls, celda, negrita=False, tam=8, relleno=None, borde=True, h='center', v=None,
                ajuste=False, fuente=None, color_borde=None):
        celda.font = Font(name=fuente or cls.FUENTE, size=tam, bold=negrita, color='FF000000')
        if relleno:
            celda.fill = PatternFill(fill_type='solid', fgColor=relleno)
        if borde:
            lado = Side(style='thin', color=color_borde or cls.BORDE)
            celda.border = Border(left=lado, right=lado, top=lado, bottom=lado)
        celda.alignment = Alignment(horizontal=h, vertical=v, wrap_text=ajuste)

    @classmethod
    def _escribir(cls, ws, fila, col, valor=None, **estilo):
        celda = ws.cell(row=fila, column=col)
        if valor is not None:
            celda.value = valor
        cls._estilo(celda, **estilo)
        return celda

    @classmethod
    def _combinar(cls, ws, fila, col, fila_fin, col_fin, valor=None, **estilo):
        for f in range(fila, fila_fin + 1):
            for c in range(col, col_fin + 1):
                cls._estilo(ws.cell(row=f, column=c), **estilo)
        if valor is not None:
            ws.cell(row=fila, column=col).value = valor
        if (fila, col) != (fila_fin, col_fin):
            ws.merge_cells(start_row=fila, start_column=col, end_row=fila_fin, end_column=col_fin)

    @classmethod
    def _formato_hoja(cls, ws, anchos=None):
        ws.sheet_format.defaultRowHeight = 14.25
        ws.sheet_format.customHeight = True
        for col, ancho in (anchos or cls.ANCHOS).items():
            ws.column_dimensions[col].width = ancho
        ws.page_margins = PageMargins(left=0.63, right=0.63, top=0.94, bottom=0.94)

    @classmethod
    def _encabezado(cls, ws, filas):
        """Bloque superior de cada hoja (filas 2 a 4) y el año en B6, como en el Excel 2026."""
        for i in range(3):
            fila = 2 + i
            rotulo, valor = filas[i] if i < len(filas) else (None, None)
            cls._combinar(ws, fila, 4, fila, 5, rotulo, negrita=True, borde=False, v='center')
            cls._combinar(ws, fila, 6, fila, 8, valor, negrita=True, tam=9, borde=False)
        cls._escribir(ws, 6, 2, DatosColegio.ANIO, negrita=True, tam=11, borde=False, fuente='Calibri')
        ws.row_dimensions[7].height = 5.25

    @classmethod
    def _grilla(cls, ws, fila0, col0, max_b, contenido):
        """
        Dibuja una grilla HORAS x LUNES-VIERNES con sus recreos (y la colación si llega al
        bloque 10). contenido(dia, bloque) -> (texto, relleno); en la colación bloque = 'C'.
        Retorna la última fila utilizada.
        """
        cls._combinar(ws, fila0, col0, fila0, col0 + 1, 'HORAS', negrita=True, relleno=cls.AMARILLO)
        for i, dia in enumerate(DatosColegio.DIAS):
            cls._escribir(ws, fila0, col0 + 2 + i, dia.upper(), negrita=True, relleno=cls.AMARILLO)
        fila = fila0
        for b in range(1, max_b + 1):
            fila += 1
            inicio, fin = DatosColegio.HORARIOS_BLOQUES[b]
            cls._escribir(ws, fila, col0, b, negrita=True, relleno=cls.CREMA)
            cls._escribir(ws, fila, col0 + 1, f'{inicio} - {fin}', negrita=True, relleno=cls.CREMA)
            for i, dia in enumerate(DatosColegio.DIAS):
                texto, relleno = contenido(dia, b)
                cls._escribir(ws, fila, col0 + 2 + i, texto, relleno=relleno or cls.BLANCO)
            if b in DatosColegio.RECREOS and b < max_b:
                fila += 1
                nombre, rango = DatosColegio.RECREOS[b]
                cls._escribir(ws, fila, col0, nombre, negrita=True, relleno=cls.AMARILLO)
                cls._escribir(ws, fila, col0 + 1, rango, negrita=True, relleno=cls.AMARILLO)
                for i, dia in enumerate(DatosColegio.DIAS):
                    texto, relleno = contenido(dia, 'C') if nombre == 'COLACIÓN' else (None, None)
                    cls._escribir(ws, fila, col0 + 2 + i, texto, relleno=relleno or cls.AMARILLO)
        return fila

    @classmethod
    def _tabla_horas(cls, ws, fila0, filas, rotulo_tercera='PROFESOR/A', color_etiquetas=None):
        """
        Tabla de horas bajo la grilla: ASIGNATURAS (B:D) | N° DE HORAS (E) | PROFESOR/A (F:G).
        filas: (tipo, etiqueta, horas, texto) con tipo 'encabezado', 'fila', 'subtitulo', 'artes',
        'musica' o 'total'. Retorna la fila del total.
        """
        fila = fila0
        for tipo, etiqueta, horas, texto in filas:
            if tipo == 'total':
                for col in (2, 3, 4, 6, 7):
                    cls._escribir(ws, fila, col, relleno=cls.BLANCO, borde=False, h=None, ajuste=True)
                cls._escribir(ws, fila, 5, horas, relleno=cls.AMARILLO, ajuste=True)
                return fila
            largo = len(str(etiqueta)) > 30 or len(str(texto or '')) > 26
            ws.row_dimensions[fila].height = 28.5 if largo else 14.25
            cls._combinar(ws, fila, 2, fila, 4, etiqueta, negrita=tipo != 'subtitulo', relleno=cls.AMARILLO,
                          h='left' if tipo == 'subtitulo' else None, ajuste=True, color_borde=color_etiquetas)
            if tipo == 'encabezado':
                cls._escribir(ws, fila, 5, 'N° DE HORAS', negrita=True, relleno=cls.AMARILLO, ajuste=True)
                cls._combinar(ws, fila, 6, fila, 7, rotulo_tercera, negrita=True, relleno=cls.AMARILLO, ajuste=True)
            else:
                if tipo == 'artes':
                    cls._combinar(ws, fila, 5, fila + 1, 5, horas, relleno=cls.BLANCO, v='center', ajuste=True)
                elif tipo != 'musica':
                    cls._escribir(ws, fila, 5, horas, relleno=cls.BLANCO, v='center', ajuste=True)
                cls._combinar(ws, fila, 6, fila, 7, texto, relleno=cls.BLANCO, ajuste=True)
            fila += 1
        return fila - 1

    @classmethod
    def _leyenda_skills(cls, ws, fila):
        cls._escribir(ws, fila, 2, None, relleno=cls.SKILLS, borde=False)
        cls._escribir(ws, fila, 3, 'English skills', negrita=True, tam=10, borde=False, h=None, fuente='Arial')

    @staticmethod
    def _nombre_hoja(nombre):
        for ch in '[]:*?/\\':
            nombre = nombre.replace(ch, '')
        return nombre[:31]

    # ------------------------------------------------------------------
    # Libro de cursos
    # ------------------------------------------------------------------
    @classmethod
    def exportar_horario_cursos(cls, horario, ruta_salida):
        """
        Crea el libro de cursos con la estructura y el diseño de HORARIO CURSOS 2026.xlsx:
        - Hoja 'SALAS ELECTIVOS' (English skills, electivos, artes, música y pastoral).
        - Hoja 'GIMNASIOS' (Gimnasio A, Gimnasio B y Patio Santo Domingo).
        - Hojas de párvulos (Ed. Física) y una hoja por cada uno de los 24 cursos.
        - Hoja 'AUDITORIA' con la verificación de restricciones.
        """
        wb = cls._nuevo_libro()
        os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
        ws = wb.active
        ws.title = 'SALAS ELECTIVOS'
        cls._hoja_salas(ws, horario)
        cls._hoja_gimnasios(wb.create_sheet('GIMNASIOS'), horario)
        for curso in DatosColegio.todos_los_cursos():
            cls._hoja_curso(wb.create_sheet(curso), curso, horario)
        cls._hoja_auditoria(wb.create_sheet('AUDITORIA'), horario)
        wb.save(ruta_salida)
        return ruta_salida

    # Alias por retrocompatibilidad
    exportar_horario = exportar_horario_cursos

    @classmethod
    def _hoja_curso(cls, ws, curso, horario):
        cls._formato_hoja(ws)
        cls._insertar_escudo(ws)
        cls._encabezado(ws, [('CURSO', curso), ('PROFESOR/A JEFE', DatosColegio.PROFESORES_JEFES.get(curso, ''))])
        parvulo = DatosColegio.es_parvulo(curso)
        max_b = 10 if parvulo or max(DatosColegio.JORNADAS[curso]) > 8 else 8
        bloqueos = DatosColegio.BLOQUEOS_PASTORAL['PÁRVULOS'] if parvulo else []

        def contenido(dia, b):
            if (dia, b) in bloqueos:
                return 'PASTORAL', cls.PASTORAL
            if b == 'C':
                return None, None
            if parvulo and b not in DatosColegio.get_bloques_permitidos(curso, dia):
                return None, None   # Como en el Excel 2026: el resto de la jornada parvularia queda en blanco
            if b not in DatosColegio.get_bloques_permitidos(curso, dia):
                return None, (cls.AMARILLO if b > 8 else cls.CREMA)
            item = horario.asignaciones[curso][dia].get(b)
            if not item:
                return None, None
            return item['etiqueta'], (cls.SKILLS if item['asignatura'] == 'English Skills' else None)

        ultima = cls._grilla(ws, 8, 2, max_b, contenido)
        color_etiquetas = cls.BORDE_MEDIA if DatosColegio.nivel(curso) in DatosColegio.FRANJAS_ELECTIVOS else None
        fila_total = cls._tabla_horas(ws, ultima + (3 if max_b == 8 else 2), cls._filas_tabla_curso(curso),
                                      color_etiquetas=color_etiquetas)
        if any(a == 'English Skills' for a, _, _, _ in DatosColegio.get_malla_curricular()[curso]):
            cls._leyenda_skills(ws, fila_total + 2)

    @classmethod
    def _docentes_electivo(cls, nivel, electivo):
        por_docente = {}
        for franja in DatosColegio.FRANJAS_ELECTIVOS[nivel]:
            for asig, seccion, doc, _, _ in franja['grupos']:
                if asig == electivo:
                    por_docente.setdefault(doc, []).append(seccion)
        return ' / '.join(f"{d} ({', '.join(s)})" for d, s in por_docente.items())

    @classmethod
    def _filas_tabla_curso(cls, curso):
        """Filas de la tabla de horas de un curso, con las mismas asignaturas que el Excel 2026."""
        malla = DatosColegio.get_malla_curricular()[curso]
        lecciones = {a: (h, DatosColegio.docentes_de(d)) for a, h, d, _ in malla}
        total = sum(h for _, h, _, _ in malla)

        def horas(*asigs):
            return sum(lecciones[a][0] for a in asigs if a in lecciones)

        def profes(*asigs):
            docs = []
            for a in asigs:
                for d in lecciones.get(a, (0, ()))[1]:
                    if d not in docs:
                        docs.append(d)
            return ' / '.join(docs)

        def fila(etiqueta, *asigs):
            return ('fila', etiqueta, horas(*asigs), profes(*asigs))

        encabezado = ('encabezado', 'ASIGNATURAS', None, None)
        if DatosColegio.es_parvulo(curso):
            return [encabezado, fila('EDUCACIÓN FÍSICA', 'Educación Física y Salud'), ('total', None, total, None)]

        ingles = profes('Idioma Extranjero Inglés')
        if 'English Skills' in lecciones:
            ingles += f" / {lecciones['English Skills'][1][1]} (E. skills)"
        fila_ingles = ('fila', 'IDIOMA EXTRANJERO: INGLÉS', horas('Idioma Extranjero Inglés', 'English Skills'), ingles)
        numero = int(curso[0])

        if 'BÁSICO' in curso and numero <= 4:
            orientec = profes('Orientación / Tecnología')
            media_hora = horas('Orientación / Tecnología') / 2
            filas = [
                encabezado,
                fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                fila_ingles,
                fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Ciencias Sociales'),
                fila('CIENCIAS NATURALES', 'Ciencias Naturales'),
                ('fila', 'EDUCACIÓN TECNOLÓGICA', media_hora, orientec),
                fila('ARTES VISUALES', 'Artes Visuales'),
                fila('MÚSICA', 'Educación Musical'),
                fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                ('fila', 'ORIENTACIÓN', media_hora, orientec),
                fila('RELIGIÓN', 'Religión'),
            ]
        elif 'BÁSICO' in curso:
            filas = [
                encabezado,
                fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                fila_ingles,
                fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Historia, Geografía y CC.SS.'),
                fila('CIENCIAS NATURALES', 'Ciencias Naturales'),
                fila('EDUCACIÓN TECNOLÓGICA', 'Educación Tecnológica'),
            ]
            if numero >= 7:
                filas.append(fila('FÍSICA', 'Física'))
            filas += [
                fila('ARTES VISUALES', 'Artes Visuales'),
                fila('MÚSICA', 'Educación Musical'),
                fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                fila('ORIENTACIÓN', 'Orientación'),
                fila('RELIGIÓN', 'Religión'),
            ]
        else:
            artes, musica = lecciones['Artes Visuales / Música'][1]
            artes_musica = [('artes', 'ARTES VISUALES', horas('Artes Visuales / Música'), artes),
                            ('musica', 'MÚSICA', None, musica)]
            cierre = [fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                      fila('ORIENTACIÓN', 'Orientación'),
                      fila('RELIGIÓN', 'Religión')]
            if numero <= 2:
                filas = [
                    encabezado,
                    fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                    fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                    fila_ingles,
                    fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Historia, Geografía y CC.SS.'),
                    ('subtitulo', 'CIENCIAS NATURALES:', None, None),
                    fila('BIOLOGÍA', 'Biología'),
                    fila('QUÍMICA', 'Química'),
                    fila('FÍSICA', 'Física'),
                    fila('EDUCACIÓN TECNOLÓGICA', 'Educación Tecnológica'),
                ] + artes_musica + cierre
            else:
                nivel = DatosColegio.nivel(curso)
                filas = [
                    ('encabezado', 'ASIGNATURAS - FORM. COMUN - LIBRE DISP.', None, None),
                    fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                    fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                    fila_ingles,
                    fila('EDUCACIÓN CIUDADANA', 'Educación Ciudadana'),
                    fila('CIENCIAS PARA LA CIUDADANIA', 'Ciencias para la Ciudadanía'),
                    fila('FILOSOFÍA', 'Filosofía'),
                    fila('FÍSICA', 'Física'),
                ] + artes_musica + cierre
                filas.append(('encabezado', 'ASIGNATURAS - FORMACIÓN DIFERENCIADA', None, None))
                for electivo in DatosColegio.ELECTIVOS_OFERTA[nivel]:
                    filas.append(('fila', electivo, 6, cls._docentes_electivo(nivel, electivo)))
        filas.append(('total', None, total, None))
        return filas

    @classmethod
    def _hoja_gimnasios(cls, ws, horario):
        cls._formato_hoja(ws)
        recintos = horario.asignar_recintos()
        for rango in ('D1:E1', 'F1:H1', 'F2:H2', 'D3:E3', 'F3:H3', 'D19:E19', 'F18:H18', 'F19:H19',
                      'D34:E34', 'F33:H33', 'F34:H34'):
            ws.merge_cells(rango)
        for recinto, fila_titulo, fila_grilla in (('GIMNASIO A', 2, 5), ('GIMNASIO B', 18, 20),
                                                  ('PATIO SANTO DOMINGO', 33, 35)):
            cls._combinar(ws, fila_titulo, 4, fila_titulo, 5, recinto, negrita=True, tam=10, borde=False, v='center')

            def contenido(dia, b, recinto=recinto):
                if b == 'C':
                    return None, None
                cursos = recintos.get((recinto, dia, b))
                if cursos:
                    return ' / '.join(DatosColegio.codigo_corto(c) for c in cursos), None
                # Los pares libres de los gimnasios se marcan "GIMNASIO / DISPONIBLE", como en el Excel 2026
                b1 = b if b % 2 else b - 1
                if recinto != 'PATIO SANTO DOMINGO' and not recintos.get((recinto, dia, b1)) \
                        and not recintos.get((recinto, dia, b1 + 1)):
                    return ('GIMNASIO' if b % 2 else 'DISPONIBLE'), None
                return None, None

            cls._grilla(ws, fila_grilla, 2, 8, contenido)
        for fila, alto in {4: 5.25, 17: 8.25, 19: 1.5, 32: 6.75, 34: 2.25}.items():
            ws.row_dimensions[fila].height = alto

    @classmethod
    def _hoja_salas(cls, ws, horario):
        anchos = {'A': 0.66, 'I': 10.66, 'Q': 10.66}
        for base in (2, 10, 18):
            anchos[get_column_letter(base)] = 8.66
            anchos[get_column_letter(base + 1)] = 12.66
            for i in range(2, 7):
                anchos[get_column_letter(base + i)] = 13.89
        cls._formato_hoja(ws, anchos)

        contenido_salas = defaultdict(list)   # (sala, dia, bloque) -> [textos]
        for curso in DatosColegio.CURSOS:
            for dia in DatosColegio.DIAS:
                for b, item in horario.asignaciones[curso][dia].items():
                    if item['asignatura'] == 'English Skills':
                        contenido_salas[('ELECTIVO 1', dia, b)].append(DatosColegio.codigo_corto(curso))
                    elif item['asignatura'] == 'Artes Visuales / Música':
                        contenido_salas[('SALA DE ARTES', dia, b)].append(DatosColegio.codigo_sala(curso))
                        contenido_salas[('SALA DE MÚSICA', dia, b)].append(DatosColegio.codigo_sala(curso))
                    elif item['espacio'] == 'Sala de Música':
                        contenido_salas[('SALA DE MÚSICA', dia, b)].append(DatosColegio.codigo_sala(curso))
        for franjas in DatosColegio.FRANJAS_ELECTIVOS.values():
            for franja in franjas:
                for dia, b0 in franja['periodos']:
                    for b in (b0, b0 + 1):
                        for _, _, _, sala, etiqueta in franja['grupos']:
                            contenido_salas[(sala, dia, b)].append(etiqueta)

        for fila_titulo, col0, sala, subtitulo in cls.DISPOSICION_SALAS:
            cls._combinar(ws, fila_titulo, col0 + 2, fila_titulo, col0 + 3, sala, negrita=True, tam=11, borde=False)
            cls._combinar(ws, fila_titulo, col0 + 4, fila_titulo, col0 + 6, subtitulo, negrita=True, borde=False)
            ws.row_dimensions[fila_titulo + 2].height = 5.25
            bloqueos = DatosColegio.BLOQUEOS_PASTORAL.get(sala, [])

            def contenido(dia, b, sala=sala, bloqueos=bloqueos):
                if (dia, b) in bloqueos:
                    return 'PASTORAL', cls.PASTORAL
                textos = contenido_salas.get((sala, dia, b))
                return (' / '.join(textos) if textos else None), None

            cls._grilla(ws, fila_titulo + 3, col0, 10, contenido)

    @classmethod
    def _hoja_auditoria(cls, ws, horario):
        auditoria = ValidadorRestricciones.auditar(horario)
        m = auditoria['metricas']
        cls._formato_hoja(ws, {'A': 0.66, 'B': 46, 'C': 42, 'D': 22})
        cls._escribir(ws, 2, 2, f'INFORME DE CUMPLIMIENTO Y FACTIBILIDAD — AÑO ESCOLAR {DatosColegio.ANIO}',
                      negrita=True, tam=11, borde=False, h=None)
        estado = ('FACTIBLE (100% restricciones duras cumplidas)' if auditoria['valido']
                  else f"NO FACTIBLE ({len(auditoria['errores_duros'])} errores)")
        cls._escribir(ws, 3, 2, 'Estado global:', negrita=True, borde=False, h=None)
        cls._escribir(ws, 3, 3, estado, negrita=True, borde=False, h=None)

        filas = [
            ('Colisiones docentes (no solapamiento)', f"{m['colisiones_docentes']} colisiones", m['colisiones_docentes']),
            ('Cargas curriculares y jornadas completas',
             f"{m['total_bloques_asignados']} / {m['total_bloques_esperados']} bloques", m['errores_carga']),
            ('Orientación a cargo del profesor jefe', f"{m['errores_jefatura']} inconsistencias", m['errores_jefatura']),
            ('Sincronización de electivos 3° y 4° medio', f"{m['errores_electivos']} diferencias A/B", m['errores_electivos']),
            ('Recintos deportivos (1 curso por recinto)', f"{m['errores_recintos']} excesos", m['errores_recintos']),
            ('Sala English skills y bloqueos de pastoral', f"{m['errores_salas']} conflictos", m['errores_salas']),
            ('Una sesión diaria por asignatura', f"{m['mismo_dia_violaciones']} repeticiones", m['mismo_dia_violaciones']),
            ('Troncales en bloques consecutivos', f"{m['core_no_consecutivo']} bloques separados", m['core_no_consecutivo']),
            ('Bloques dobles (90 min)', f"{m['porcentaje_bloques_dobles']}% de las horas", None),
            ('Ventanas docentes', f"{m['total_ventanas_docentes']} bloques de espera", None),
            ('Uso del Patio Santo Domingo', f"{m['uso_patio']} bloques", None),
        ]
        for i, texto in enumerate(['MÉTRICA / CRITERIO', 'VALOR EVALUADO', 'ESTADO']):
            cls._escribir(ws, 5, 2 + i, texto, negrita=True, relleno=cls.AMARILLO)
        fila = 6
        for criterio, valor, errores in filas:
            estado = 'INDICADOR' if errores is None else ('CUMPLE' if errores == 0 else 'NO CUMPLE')
            cls._escribir(ws, fila, 2, criterio, negrita=True, relleno=cls.CREMA, h=None)
            cls._escribir(ws, fila, 3, valor, relleno=cls.BLANCO)
            cls._escribir(ws, fila, 4, estado, relleno=cls.BLANCO)
            fila += 1
        if auditoria['errores_duros']:
            fila += 1
            cls._escribir(ws, fila, 2, 'DETALLE DE ERRORES', negrita=True, relleno=cls.AMARILLO, h=None)
            for error in auditoria['errores_duros'][:300]:
                fila += 1
                cls._combinar(ws, fila, 2, fila, 4, error, relleno=cls.BLANCO, h=None)

    # ------------------------------------------------------------------
    # Libro de docentes
    # ------------------------------------------------------------------
    @classmethod
    def exportar_horario_docentes(cls, horario, ruta_salida):
        """
        Crea el libro de docentes:
        - Hoja 'DISTRIBUCIÓN HORARIA' con el formato de DISTRIBUCIÓN HORARIA 2026.docx
          (una tabla por departamento: profesor, cursos, n° de horas, asignatura y horas lectivas).
        - Una hoja por docente con su grilla semanal y el mismo diseño de las hojas de curso.
        """
        wb = cls._nuevo_libro()
        os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
        ws = wb.active
        ws.title = 'DISTRIBUCIÓN HORARIA'
        cls._hoja_distribucion(ws, horario)
        for depto, docentes in DatosColegio.DEPARTAMENTOS:
            for doc in docentes:
                cls._hoja_docente(wb.create_sheet(cls._nombre_hoja(doc)), doc, depto, horario)
        wb.save(ruta_salida)
        return ruta_salida

    @classmethod
    def _resumen_docente(cls, doc, horario):
        """Filas de la distribución horaria de un docente, calculadas desde el horario generado."""
        orden = {c: i for i, c in enumerate(DatosColegio.todos_los_cursos())}
        normales = defaultdict(int)
        electivos = defaultdict(list)
        for (dia, b), (curso, asig) in horario.docente_ocupado.get(doc, {}).items():
            if curso in orden:
                normales[(curso, asig)] += 1
            else:
                nombre, seccion = asig.rsplit(' (', 1)
                electivos[(curso, nombre)].append(seccion.rstrip(')'))

        filas = []
        for curso, asig in sorted(normales, key=lambda k: (k[1] == 'Orientación', orden[k[0]])):
            if asig == 'Orientación':
                nombre, tipo = 'Jefatura', 'jefatura'
            elif asig == 'English Skills':
                nombre, tipo = 'Inglés (English skills)', 'skills'
            elif asig == 'Orientación / Tecnología':
                nombre, tipo = 'Tecn/C.Orien.', 'orientec'
            elif asig == 'Artes Visuales / Música':
                nombre, tipo = ('Música' if doc == 'Música' else 'Artes'), 'normal'
            else:
                nombre, tipo = cls.NOMBRES_DISTRIBUCION.get(asig, asig), 'normal'
            filas.append({'cursos': DatosColegio.codigo_corto(curso), 'curso': curso,
                          'horas': normales[(curso, asig)], 'asignatura': nombre, 'tipo': tipo})
        for (grupo, nombre), secciones in electivos.items():
            unicas = list(dict.fromkeys(secciones))
            cursos = ', '.join(unicas)
            if all(s.startswith('S.') for s in unicas):   # Secciones mixtas A+B: "3°AB S.1, S.2"
                cursos = f"{grupo.replace(' MEDIO A-B', 'AB')} {cursos}"
            posicion = sum(1 for f in filas if f['tipo'] != 'jefatura')
            filas.insert(posicion, {'cursos': cursos,
                                    'curso': f"{grupo} ({', '.join(unicas)})",
                                    'horas': len(secciones), 'asignatura': nombre, 'tipo': 'electivo'})
        for cursos, actividad, horas in DatosColegio.HORAS_NO_LECTIVAS.get(doc, []):
            filas.append({'cursos': cursos, 'curso': cursos or 'No programada en grilla', 'horas': horas,
                          'asignatura': actividad, 'tipo': 'no_lectiva'})
        return filas, sum(f['horas'] for f in filas)

    @classmethod
    def _hoja_distribucion(cls, ws, horario):
        for col, ancho in {'A': 26, 'B': 20, 'C': 13, 'D': 34, 'E': 16}.items():
            ws.column_dimensions[col].width = ancho
        ws.page_margins = PageMargins(left=0.63, right=0.63, top=0.94, bottom=0.94)
        cls._combinar(ws, 1, 1, 1, 5, f'DISTRIBUCIÓN DE HORAS POR ASIGNATURAS – AÑO ACADÉMICO {DatosColegio.ANIO}.',
                      negrita=True, tam=14, borde=False, h=None, fuente='Calibri')
        rellenos = {'jefatura': cls.CELESTE, 'skills': cls.VERDE, 'orientec': cls.AZUL}
        estilo = {'tam': 10, 'fuente': 'Abadi', 'v': 'center', 'color_borde': cls.BORDE_DOCUMENTO}
        fila = 3
        for titulo, docentes in DatosColegio.DEPARTAMENTOS:
            cls._combinar(ws, fila, 1, fila, 5, f'ASIGNATURAS: {titulo}', negrita=True, tam=12, borde=False,
                          h=None, ajuste=True, fuente='Calibri')
            ws.row_dimensions[fila].height = 32 if len(titulo) > 60 else 18
            fila += 1
            if titulo == 'INGLÉS':
                cls._combinar(ws, fila, 1, fila, 5, 'Verdes: English skills (2 docentes en el mismo bloque)',
                              tam=10, borde=False, h=None, fuente='Calibri')
                fila += 1
            for i, texto in enumerate(['PROFESOR', 'CURSOS', 'Nº DE HORAS', 'ASIGNATURA', 'HORAS LECTIVAS'], 1):
                cls._escribir(ws, fila, i, texto, negrita=True, relleno=cls.GRIS, ajuste=True, **estilo)
            fila += 1
            for doc in docentes:
                filas, total = cls._resumen_docente(doc, horario)
                if not filas:
                    continue
                inicio = fila
                for f in filas:
                    relleno = rellenos.get(f['tipo'])
                    cls._escribir(ws, fila, 2, f['cursos'], relleno=relleno, **estilo)
                    cls._escribir(ws, fila, 3, f['horas'], relleno=relleno, **estilo)
                    cls._escribir(ws, fila, 4, f['asignatura'], h=None, **estilo)
                    fila += 1
                cls._combinar(ws, inicio, 1, fila - 1, 1, doc, **estilo)
                cls._combinar(ws, inicio, 5, fila - 1, 5, total, **estilo)
            fila += 1

    @classmethod
    def _hoja_docente(cls, ws, doc, depto, horario):
        anchos = dict(cls.ANCHOS)
        anchos.update({col: 17 for col in 'DEFGH'})
        cls._formato_hoja(ws, anchos)
        cls._insertar_escudo(ws)
        ocupado = horario.docente_ocupado.get(doc, {})
        detalle = horario.docente_detalle.get(doc, {})
        filas, total = cls._resumen_docente(doc, horario)
        horas = total if total == len(ocupado) else f"{total} ({len(ocupado)} en grilla)"
        departamento = depto.split(' –')[0].split(' - ')[0]
        cls._encabezado(ws, [('DOCENTE', doc), ('DEPARTAMENTO', departamento), ('HORAS LECTIVAS', horas)])

        def contenido(dia, b):
            if b == 'C' or (dia, b) not in ocupado:
                return None, None
            if (dia, b) in detalle:
                return detalle[(dia, b)]['etiqueta'], None
            curso, asig = ocupado[(dia, b)]
            etiqueta = horario.asignaciones[curso][dia][b]['etiqueta']
            if asig == 'Artes Visuales / Música':
                etiqueta = 'MÚSICA' if doc == 'Música' else 'ARTES'
            return f"{DatosColegio.codigo_corto(curso)} {etiqueta}", (cls.SKILLS if asig == 'English Skills' else None)

        max_b = 10 if any(b > 8 for (_, b) in ocupado) else 8
        ultima = cls._grilla(ws, 8, 2, max_b, contenido)
        tabla = [('encabezado', 'ASIGNATURAS', None, None)]
        for f in filas:
            tabla.append(('fila', f['asignatura'].upper(), f['horas'], f['curso']))
        tabla.append(('total', None, total, None))
        fila_total = cls._tabla_horas(ws, ultima + (3 if max_b == 8 else 2), tabla, rotulo_tercera='CURSO')
        if any(f['tipo'] == 'skills' for f in filas):
            cls._leyenda_skills(ws, fila_total + 2)


# =============================================================================
# 6. MENÚ INTERACTIVO Y VISTA POR CONSOLA
# =============================================================================

class MenuInteractivo:
    """Consola interactiva para explorar los horarios generados y exportar."""

    def __init__(self, horario):
        self.horario = horario
        carpeta_outputs = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Outputs Excel")
        self.ruta_excel_cursos = os.path.join(carpeta_outputs, "Horario_Cursos_MMDD.xlsx")
        self.ruta_excel_docentes = os.path.join(carpeta_outputs, "Horarios_Docentes_Colegio_MMDD.xlsx")
        self.ruta_excel = self.ruta_excel_cursos

    @staticmethod
    def _recortar(texto, ancho):
        texto = str(texto)
        return texto if len(texto) <= ancho else texto[:ancho - 2] + '..'

    @staticmethod
    def _imprimir_recreo(b, col_w, header):
        nombre, rango = DatosColegio.RECREOS[b]
        barra = f"--- {nombre} {rango} ---".center(5 * (col_w + 3) - 3, '-')
        print(f"{'':<6} | {'':<13} | {barra}")
        print("-" * len(header))

    def mostrar_tabla_curso(self, curso):
        """Imprime en terminal el horario semanal formateado para un curso."""
        curso = curso.upper().strip()
        if curso not in DatosColegio.todos_los_cursos():
            print(f"\n❌ Error: El curso '{curso}' no existe.")
            return

        jefe = DatosColegio.PROFESORES_JEFES.get(curso, 'No asignado')
        col_w = 17
        ancho_total = 6 + 3 + 13 + 3 + 5 * (col_w + 3) - 1
        print("\n" + "=" * ancho_total)
        print(f"  COLEGIO MADRES DOMINICAS — HORARIO SEMANAL {DatosColegio.ANIO}: {curso}")
        print(f"  Profesor(a) Jefe: {jefe}")
        print("=" * ancho_total)

        header = f"{'BLOQUE':<6} | {'HORARIO':<13} | " + " | ".join(f"{d:<{col_w}}" for d in DatosColegio.DIAS)
        print(header)
        print("-" * len(header))

        max_bloques = 10 if max(DatosColegio.JORNADAS[curso]) > 8 else 8
        for b in range(1, max_bloques + 1):
            h_in, h_fi = DatosColegio.HORARIOS_BLOQUES[b]
            cols_asig = []
            cols_doc = []
            for d in DatosColegio.DIAS:
                if b not in DatosColegio.get_bloques_permitidos(curso, d):
                    cols_asig.append(f"{'---':<{col_w}}")
                    cols_doc.append(f"{' ':<{col_w}}")
                elif b in self.horario.asignaciones[curso][d]:
                    info = self.horario.asignaciones[curso][d][b]
                    docente = info['docente']
                    if info['asignatura'].startswith('Formación Diferenciada'):
                        docente = f"{len(info['docentes'])} docentes"
                    cols_asig.append(f"{self._recortar(info['etiqueta'], col_w):<{col_w}}")
                    cols_doc.append(f"{self._recortar('(' + docente + ')', col_w):<{col_w}}")
                else:
                    cols_asig.append(f"{'Libre':<{col_w}}")
                    cols_doc.append(f"{' ':<{col_w}}")

            print(f"{'B.' + str(b):<6} | {h_in + '-' + h_fi:<13} | " + " | ".join(cols_asig))
            print(f"{'':<6} | {'':<13} | " + " | ".join(cols_doc))
            if b in DatosColegio.RECREOS and b < max_bloques:
                self._imprimir_recreo(b, col_w, header)
            else:
                print("-" * len(header))

        print("=" * ancho_total)

    def mostrar_tabla_docente(self, docente_nombre):
        """Imprime la malla semanal de un docente específico."""
        coincidencias = sorted(d for d in self.horario.docente_ocupado if docente_nombre.lower() in d.lower())
        if not coincidencias:
            print(f"\n❌ No se encontró ningún docente con el nombre '{docente_nombre}'.")
            return

        exactos = [d for d in coincidencias if d.lower() == docente_nombre.lower()]
        doc = (exactos or coincidencias)[0]
        ocupado = self.horario.docente_ocupado[doc]
        detalle = self.horario.docente_detalle.get(doc, {})
        col_w = 17
        ancho_total = 6 + 3 + 13 + 3 + 5 * (col_w + 3) - 1
        print("\n" + "=" * ancho_total)
        print(f"  MALLA SEMANAL {DatosColegio.ANIO}: {doc} ({len(ocupado)} horas en grilla)")
        print("=" * ancho_total)

        header = f"{'BLOQUE':<6} | {'HORARIO':<13} | " + " | ".join(f"{d:<{col_w}}" for d in DatosColegio.DIAS)
        print(header)
        print("-" * len(header))

        max_bloques = 10 if any(b > 8 for (_, b) in ocupado) else 8
        for b in range(1, max_bloques + 1):
            h_in, h_fi = DatosColegio.HORARIOS_BLOQUES[b]
            cols_curso = []
            cols_asig = []
            for d in DatosColegio.DIAS:
                if (d, b) in ocupado:
                    c, asig = ocupado[(d, b)]
                    etiqueta = detalle[(d, b)]['etiqueta'] if (d, b) in detalle else DatosColegio.ETIQUETAS.get(asig, asig)
                    cols_curso.append(f"{self._recortar(c, col_w):<{col_w}}")
                    cols_asig.append(f"{self._recortar('(' + etiqueta + ')', col_w):<{col_w}}")
                else:
                    cols_curso.append(f"{'Libre':<{col_w}}")
                    cols_asig.append(f"{' ':<{col_w}}")

            print(f"{'B.' + str(b):<6} | {h_in + '-' + h_fi:<13} | " + " | ".join(cols_curso))
            print(f"{'':<6} | {'':<13} | " + " | ".join(cols_asig))
            if b in DatosColegio.RECREOS and b < max_bloques:
                self._imprimir_recreo(b, col_w, header)
            else:
                print("-" * len(header))
        print("=" * ancho_total)

    def exportar(self):
        """Genera ambos libros Excel en la carpeta Outputs Excel/."""
        print(f"\n⏳ Generando libros Excel en carpeta 'Outputs Excel/' ...")

        # 1. Horario de cursos
        ExportadorExcel.exportar_horario_cursos(self.horario, self.ruta_excel_cursos)
        print(f"✅ ¡Horario de Cursos generado exitosamente!")
        print(f"📁 [1] Cursos:   {self.ruta_excel_cursos}")

        # 2. Horario de docentes
        ExportadorExcel.exportar_horario_docentes(self.horario, self.ruta_excel_docentes)
        print(f"✅ ¡Horarios de Docentes generado exitosamente!")
        print(f"📁 [2] Docentes: {self.ruta_excel_docentes}\n")

    def mostrar_auditoria(self):
        """Muestra el reporte de verificación de restricciones."""
        auditoria = ValidadorRestricciones.auditar(self.horario)
        m = auditoria['metricas']
        print("\n" + "=" * 70)
        print(f"  AUDITORÍA DE RESTRICCIONES {DatosColegio.ANIO} — COLEGIO MADRES DOMINICAS")
        print("=" * 70)
        estado = "✅ FACTIBLE (100% Restricciones Duras Cumplidas)" if auditoria['valido'] else "❌ CONFLICTOS"
        print(f"  Estado Global: {estado}")
        print("-" * 70)
        print(f"  • Cursos programados:              {m['cursos_auditados']} (24 regulares + 3 párvulos)")
        print(f"  • Total bloques asignados:         {m['total_bloques_asignados']} / {m['total_bloques_esperados']}")
        print(f"  • Colisiones docentes:             {m['colisiones_docentes']}")
        print(f"  • Errores de carga o jornada:      {m['errores_carga']}")
        print(f"  • Jefaturas inconsistentes:        {m['errores_jefatura']}")
        print(f"  • Electivos desincronizados:       {m['errores_electivos']}")
        print(f"  • Recintos deportivos excedidos:   {m['errores_recintos']} (bloques en el patio: {m['uso_patio']})")
        print(f"  • Conflictos de salas/pastoral:    {m['errores_salas']}")
        print(f"  • Repeticiones diarias:            {m['mismo_dia_violaciones']}")
        print(f"  • Troncales no consecutivas:       {m['core_no_consecutivo']}")
        print(f"  • Porcentaje de Bloques Dobles:    {m['porcentaje_bloques_dobles']}% del total")
        print(f"  • Ventanas docentes acumuladas:    {m['total_ventanas_docentes']} horas libres intermedias")
        for error in auditoria['errores_duros'][:10]:
            print(f"    - {error}")
        print("=" * 70)

    def iniciar(self):
        """Bucle principal de interacción con el usuario."""
        cursos = DatosColegio.todos_los_cursos()
        while True:
            print("\n" + "╔" + "═" * 68 + "╗")
            print("║   SISTEMA DE ASIGNACIÓN DE HORARIOS - COLEGIO MMDD (TGOP) - 2026   ║")
            print("╚" + "═" * 68 + "╝")
            print("  1. Seleccionar curso para ver horario en pantalla")
            print("  2. Ver horario de todos los cursos (secuencial)")
            print("  3. Ver horario semanal de un docente")
            print("  4. Ver reporte de auditoría y verificación de restricciones")
            print("  5. Generar y exportar archivos Excel (Cursos y Docentes)")
            print("  0. Salir")
            print("-" * 70)

            opcion = input("Seleccione una opción (0-5): ").strip()

            if opcion == '1':
                print("\nLista de Cursos Disponibles:")
                for idx, c in enumerate(cursos, 1):
                    print(f"  [{idx:2d}] {c}")
                seleccion = input(f"\nIngrese el número (1-{len(cursos)}) o nombre exacto del curso: ").strip()
                if seleccion.isdigit() and 1 <= int(seleccion) <= len(cursos):
                    self.mostrar_tabla_curso(cursos[int(seleccion) - 1])
                else:
                    self.mostrar_tabla_curso(seleccion)

            elif opcion == '2':
                for c in cursos:
                    self.mostrar_tabla_curso(c)

            elif opcion == '3':
                doc_query = input("\nIngrese el nombre del docente (ej. 'Inglés 2', 'Básica 5'): ").strip()
                self.mostrar_tabla_docente(doc_query)

            elif opcion == '4':
                self.mostrar_auditoria()

            elif opcion == '5':
                self.exportar()

            elif opcion == '0':
                print("\n👋 Saliendo del sistema de horarios. ¡Hasta pronto!\n")
                break
            else:
                print("\n⚠️ Opción no válida. Intente nuevamente.")


# =============================================================================
# 7. PUNTO DE ENTRADA PRINCIPAL
# =============================================================================

def main():
    parser = argparse.ArgumentParser(description="Generador de Horarios Escolares MMDD (año 2026)")
    parser.add_argument('--curso', type=str, help="Nombre del curso a consultar directamente")
    parser.add_argument('--docente', type=str, help="Nombre del docente a consultar directamente")
    parser.add_argument('--export', action='store_true', help="Generar Excel automáticamente y salir")
    parser.add_argument('--audit', action='store_true', help="Mostrar auditoría de restricciones y salir")
    parser.add_argument('--seed', type=int, default=3, help="Semilla del generador (default: 3)")
    args = parser.parse_args()

    # 1. Generar horario factible
    motor = MotorHorarios(seed=args.seed)
    horario = motor.generar()
    menu = MenuInteractivo(horario)

    # 2. Generar siempre el Excel en Outputs Excel automáticamente
    menu.exportar()

    # 3. Procesar modos CLI
    if args.curso:
        menu.mostrar_tabla_curso(args.curso)
    elif args.docente:
        menu.mostrar_tabla_docente(args.docente)
    elif args.audit:
        menu.mostrar_auditoria()
    elif not args.export:
        # Modo interactivo en consola
        menu.iniciar()

if __name__ == '__main__':
    main()
