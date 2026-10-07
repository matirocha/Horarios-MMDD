# -*- coding: utf-8 -*-
"""
Exportación de los horarios a libros .xlsx con el diseño oficial 2026
(mismo formato que 'data/Data 2026/HORARIO CURSOS 2026.xlsx').
"""

import io
import os
import zipfile
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.drawing.image import Image as ImagenExcel
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.indexed_list import IndexedList
    from openpyxl.worksheet.page import PageMargins
except ImportError:  # El motor y la consola funcionan sin openpyxl; solo la exportación lo requiere
    openpyxl = None

from .datos import DatosColegio
from .validador import ValidadorRestricciones

# Raíz del repositorio (notebooks/horarios/ -> notebooks/ -> raíz)
RAIZ_REPOSITORIO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


class ExportadorExcel:
    """
    Genera los libros .xlsx replicando el diseño de 'data/Data 2026/HORARIO CURSOS 2026.xlsx':
    fuente Cavolini 8, encabezados y recreos en amarillo, número de bloque y horario en crema,
    bordes finos, grilla en las columnas B-H y tabla de horas por asignatura bajo cada grilla.
    """

    FUENTE = 'Cavolini'
    AMARILLO = 'FFFFE598'   # Encabezados, recreos, colación y nombres de asignaturas
    CREMA = 'FFFEF2CB'      # Número de bloque, horario y bloques fuera de la jornada
    BLANCO = 'FFFFFFFF'     # Celdas de clases
    SKILLS = 'FFFFFF00'     # English skills
    PASTORAL = 'FFFF9900'   # Bloqueos de pastoral
    BORDE = 'FFED7D31'      # Bordes finos naranjos de todas las grillas y tablas
    BORDE_MEDIA = 'FF0070C0'  # Bordes azules de las asignaturas en la tabla de 3° y 4° medio
    BORDE_DOCUMENTO = 'FF000000'  # Bordes negros de la distribución horaria (documento Word 2026)
    GRIS = 'FFC0C0C0'       # Encabezado de la distribución horaria (documento Word 2026)
    CELESTE = 'FFB7DEE8'    # Jefaturas en la distribución horaria
    VERDE = 'FFC4D79B'      # English skills en la distribución horaria
    AZUL = 'FFBDD6EE'       # Tecn/C.Orien. en la distribución horaria

    ANCHOS = {'A': 0.66, 'B': 8.66, 'C': 12.66, 'D': 13.89, 'E': 14.44,
              'F': 14.44, 'G': 14.44, 'H': 14.44, 'I': 10.66}

    # Excel oficial del que se toma el escudo del colegio (si no existe, las hojas van sin escudo)
    RUTA_EXCEL_OFICIAL = os.path.join(RAIZ_REPOSITORIO, 'data', 'Data 2026', 'HORARIO CURSOS 2026.xlsx')
    _escudo = None

    # Disposición de la hoja SALAS ELECTIVOS: (fila del título, columna inicial, sala, subtítulo)
    DISPOSICION_SALAS = [
        (2, 2, 'ELECTIVO 1', 'INGLÉS - English skills'),
        (2, 10, 'SALA DE TECNOLOGÍA', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (2, 18, 'SALA DE 4°A', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (21, 2, 'ELECTIVO 2', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (21, 10, 'SALA PADRE CUETO', 'SALA PASTORAL'),
        (21, 18, 'SALA DE 4°B', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (40, 2, 'ELECTIVO 3', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (40, 10, 'SALA MADRE PILAR', 'SALA PASTORAL'),
        (59, 2, 'SALA DE ARTES', 'ARTES VISUALES (MEDIA)'),
        (59, 10, 'SALA DE 3°A', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
        (78, 2, 'SALA DE MÚSICA', 'MÚSICA'),
        (78, 10, 'SALA DE 3°B', 'ASIGNATURAS DE PROFUNDIZACIÓN'),
    ]

    # Nombre corto de cada asignatura en la distribución horaria (como en el documento Word 2026)
    NOMBRES_DISTRIBUCION = {
        'Lenguaje y Comunicación': 'Lenguaje',
        'Educación Matemática': 'Matemática',
        'Idioma Extranjero Inglés': 'Inglés',
        'Ciencias Sociales': 'C. sociales',
        'Historia, Geografía y CC.SS.': 'Historia',
        'Educación Tecnológica': 'Tecnología',
        'Artes Visuales': 'Artes',
        'Educación Musical': 'Música',
        'Educación Física y Salud': 'Ed. Física',
        'Educación Ciudadana': 'Educación ciudadana',
        'Ciencias para la Ciudadanía': 'Ciencias para la ciudadanía',
    }

    # ------------------------------------------------------------------
    # Utilidades de formato
    # ------------------------------------------------------------------
    @classmethod
    def _nuevo_libro(cls):
        """Libro vacío cuya fuente por defecto es Cavolini 11, igual que el Excel 2026."""
        if openpyxl is None:
            raise RuntimeError("La exportación a Excel requiere el paquete openpyxl (pip install openpyxl).")
        wb = openpyxl.Workbook()
        wb._fonts = IndexedList([Font(name=cls.FUENTE, sz=11, family=2)])
        return wb

    @classmethod
    def _insertar_escudo(cls, ws):
        """Inserta el escudo del colegio en B1 (70 x 93 px), tomado del Excel oficial 2026."""
        if cls._escudo is None:
            cls._escudo = b''
            try:
                with zipfile.ZipFile(cls.RUTA_EXCEL_OFICIAL) as z:
                    medios = sorted(n for n in z.namelist() if n.startswith('xl/media/'))
                    if medios:
                        cls._escudo = z.read(medios[0])
            except (OSError, zipfile.BadZipFile):
                pass
        if not cls._escudo:
            return
        try:
            imagen = ImagenExcel(io.BytesIO(cls._escudo))
        except ImportError:   # openpyxl requiere Pillow para insertar imágenes; se omite el escudo
            return
        imagen.width, imagen.height = 70, 93
        ws.add_image(imagen, 'B1')

    @classmethod
    def _estilo(cls, celda, negrita=False, tam=8, relleno=None, borde=True, h='center', v=None,
                ajuste=False, fuente=None, color_borde=None):
        celda.font = Font(name=fuente or cls.FUENTE, size=tam, bold=negrita, color='FF000000')
        if relleno:
            celda.fill = PatternFill(fill_type='solid', fgColor=relleno)
        if borde:
            lado = Side(style='thin', color=color_borde or cls.BORDE)
            celda.border = Border(left=lado, right=lado, top=lado, bottom=lado)
        celda.alignment = Alignment(horizontal=h, vertical=v, wrap_text=ajuste)

    @classmethod
    def _escribir(cls, ws, fila, col, valor=None, **estilo):
        celda = ws.cell(row=fila, column=col)
        if valor is not None:
            celda.value = valor
        cls._estilo(celda, **estilo)
        return celda

    @classmethod
    def _combinar(cls, ws, fila, col, fila_fin, col_fin, valor=None, **estilo):
        for f in range(fila, fila_fin + 1):
            for c in range(col, col_fin + 1):
                cls._estilo(ws.cell(row=f, column=c), **estilo)
        if valor is not None:
            ws.cell(row=fila, column=col).value = valor
        if (fila, col) != (fila_fin, col_fin):
            ws.merge_cells(start_row=fila, start_column=col, end_row=fila_fin, end_column=col_fin)

    @classmethod
    def _formato_hoja(cls, ws, anchos=None):
        ws.sheet_format.defaultRowHeight = 14.25
        ws.sheet_format.customHeight = True
        for col, ancho in (anchos or cls.ANCHOS).items():
            ws.column_dimensions[col].width = ancho
        ws.page_margins = PageMargins(left=0.63, right=0.63, top=0.94, bottom=0.94)

    @classmethod
    def _encabezado(cls, ws, filas):
        """Bloque superior de cada hoja (filas 2 a 4) y el año en B6, como en el Excel 2026."""
        for i in range(3):
            fila = 2 + i
            rotulo, valor = filas[i] if i < len(filas) else (None, None)
            cls._combinar(ws, fila, 4, fila, 5, rotulo, negrita=True, borde=False, v='center')
            cls._combinar(ws, fila, 6, fila, 8, valor, negrita=True, tam=9, borde=False)
        cls._escribir(ws, 6, 2, DatosColegio.ANIO, negrita=True, tam=11, borde=False, fuente='Calibri')
        ws.row_dimensions[7].height = 5.25

    @classmethod
    def _grilla(cls, ws, fila0, col0, max_b, contenido):
        """
        Dibuja una grilla HORAS x LUNES-VIERNES con sus recreos (y la colación si llega al
        bloque 10). contenido(dia, bloque) -> (texto, relleno); en la colación bloque = 'C'.
        Retorna la última fila utilizada.
        """
        cls._combinar(ws, fila0, col0, fila0, col0 + 1, 'HORAS', negrita=True, relleno=cls.AMARILLO)
        for i, dia in enumerate(DatosColegio.DIAS):
            cls._escribir(ws, fila0, col0 + 2 + i, dia.upper(), negrita=True, relleno=cls.AMARILLO)
        fila = fila0
        for b in range(1, max_b + 1):
            fila += 1
            inicio, fin = DatosColegio.HORARIOS_BLOQUES[b]
            cls._escribir(ws, fila, col0, b, negrita=True, relleno=cls.CREMA)
            cls._escribir(ws, fila, col0 + 1, f'{inicio} - {fin}', negrita=True, relleno=cls.CREMA)
            for i, dia in enumerate(DatosColegio.DIAS):
                texto, relleno = contenido(dia, b)
                cls._escribir(ws, fila, col0 + 2 + i, texto, relleno=relleno or cls.BLANCO)
            if b in DatosColegio.RECREOS and b < max_b:
                fila += 1
                nombre, rango = DatosColegio.RECREOS[b]
                cls._escribir(ws, fila, col0, nombre, negrita=True, relleno=cls.AMARILLO)
                cls._escribir(ws, fila, col0 + 1, rango, negrita=True, relleno=cls.AMARILLO)
                for i, dia in enumerate(DatosColegio.DIAS):
                    texto, relleno = contenido(dia, 'C') if nombre == 'COLACIÓN' else (None, None)
                    cls._escribir(ws, fila, col0 + 2 + i, texto, relleno=relleno or cls.AMARILLO)
        return fila

    @classmethod
    def _tabla_horas(cls, ws, fila0, filas, rotulo_tercera='PROFESOR/A', color_etiquetas=None):
        """
        Tabla de horas bajo la grilla: ASIGNATURAS (B:D) | N° DE HORAS (E) | PROFESOR/A (F:G).
        filas: (tipo, etiqueta, horas, texto) con tipo 'encabezado', 'fila', 'subtitulo', 'artes',
        'musica' o 'total'. Retorna la fila del total.
        """
        fila = fila0
        for tipo, etiqueta, horas, texto in filas:
            if tipo == 'total':
                for col in (2, 3, 4, 6, 7):
                    cls._escribir(ws, fila, col, relleno=cls.BLANCO, borde=False, h=None, ajuste=True)
                cls._escribir(ws, fila, 5, horas, relleno=cls.AMARILLO, ajuste=True)
                return fila
            largo = len(str(etiqueta)) > 30 or len(str(texto or '')) > 26
            ws.row_dimensions[fila].height = 28.5 if largo else 14.25
            cls._combinar(ws, fila, 2, fila, 4, etiqueta, negrita=tipo != 'subtitulo', relleno=cls.AMARILLO,
                          h='left' if tipo == 'subtitulo' else None, ajuste=True, color_borde=color_etiquetas)
            if tipo == 'encabezado':
                cls._escribir(ws, fila, 5, 'N° DE HORAS', negrita=True, relleno=cls.AMARILLO, ajuste=True)
                cls._combinar(ws, fila, 6, fila, 7, rotulo_tercera, negrita=True, relleno=cls.AMARILLO, ajuste=True)
            else:
                if tipo == 'artes':
                    cls._combinar(ws, fila, 5, fila + 1, 5, horas, relleno=cls.BLANCO, v='center', ajuste=True)
                elif tipo != 'musica':
                    cls._escribir(ws, fila, 5, horas, relleno=cls.BLANCO, v='center', ajuste=True)
                cls._combinar(ws, fila, 6, fila, 7, texto, relleno=cls.BLANCO, ajuste=True)
            fila += 1
        return fila - 1

    @classmethod
    def _leyenda_skills(cls, ws, fila):
        cls._escribir(ws, fila, 2, None, relleno=cls.SKILLS, borde=False)
        cls._escribir(ws, fila, 3, 'English skills', negrita=True, tam=10, borde=False, h=None, fuente='Arial')

    @staticmethod
    def _nombre_hoja(nombre):
        for ch in '[]:*?/\\':
            nombre = nombre.replace(ch, '')
        return nombre[:31]

    # ------------------------------------------------------------------
    # Libro de cursos
    # ------------------------------------------------------------------
    @classmethod
    def exportar_horario_cursos(cls, horario, ruta_salida):
        """
        Crea el libro de cursos con la estructura y el diseño de HORARIO CURSOS 2026.xlsx:
        - Hoja 'SALAS ELECTIVOS' (English skills, electivos, artes, música y pastoral).
        - Hoja 'GIMNASIOS' (Gimnasio A, Gimnasio B y Patio Santo Domingo).
        - Hojas de párvulos (Ed. Física) y una hoja por cada uno de los 24 cursos.
        - Hoja 'AUDITORIA' con la verificación de restricciones.
        """
        wb = cls._nuevo_libro()
        os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
        ws = wb.active
        ws.title = 'SALAS ELECTIVOS'
        cls._hoja_salas(ws, horario)
        cls._hoja_gimnasios(wb.create_sheet('GIMNASIOS'), horario)
        for curso in DatosColegio.todos_los_cursos():
            cls._hoja_curso(wb.create_sheet(curso), curso, horario)
        cls._hoja_auditoria(wb.create_sheet('AUDITORIA'), horario)
        wb.save(ruta_salida)
        return ruta_salida

    # Alias por retrocompatibilidad
    exportar_horario = exportar_horario_cursos

    @classmethod
    def _hoja_curso(cls, ws, curso, horario):
        cls._formato_hoja(ws)
        cls._insertar_escudo(ws)
        cls._encabezado(ws, [('CURSO', curso), ('PROFESOR/A JEFE', DatosColegio.PROFESORES_JEFES.get(curso, ''))])
        parvulo = DatosColegio.es_parvulo(curso)
        max_b = 10 if parvulo or max(DatosColegio.JORNADAS[curso]) > 8 else 8
        bloqueos = DatosColegio.BLOQUEOS_PASTORAL['PÁRVULOS'] if parvulo else []

        def contenido(dia, b):
            if (dia, b) in bloqueos:
                return 'PASTORAL', cls.PASTORAL
            if b == 'C':
                return None, None
            if parvulo and b not in DatosColegio.get_bloques_permitidos(curso, dia):
                return None, None   # Como en el Excel 2026: el resto de la jornada parvularia queda en blanco
            if b not in DatosColegio.get_bloques_permitidos(curso, dia):
                return None, (cls.AMARILLO if b > 8 else cls.CREMA)
            item = horario.asignaciones[curso][dia].get(b)
            if not item:
                return None, None
            return item['etiqueta'], (cls.SKILLS if item['asignatura'] == 'English Skills' else None)

        ultima = cls._grilla(ws, 8, 2, max_b, contenido)
        color_etiquetas = cls.BORDE_MEDIA if DatosColegio.nivel(curso) in DatosColegio.FRANJAS_ELECTIVOS else None
        fila_total = cls._tabla_horas(ws, ultima + (3 if max_b == 8 else 2), cls._filas_tabla_curso(curso),
                                      color_etiquetas=color_etiquetas)
        if any(a == 'English Skills' for a, _, _, _ in DatosColegio.get_malla_curricular()[curso]):
            cls._leyenda_skills(ws, fila_total + 2)

    @classmethod
    def _docentes_electivo(cls, nivel, electivo):
        por_docente = {}
        for franja in DatosColegio.FRANJAS_ELECTIVOS[nivel]:
            for asig, seccion, doc, _, _ in franja['grupos']:
                if asig == electivo:
                    por_docente.setdefault(doc, []).append(seccion)
        return ' / '.join(f"{d} ({', '.join(s)})" for d, s in por_docente.items())

    @classmethod
    def _filas_tabla_curso(cls, curso):
        """Filas de la tabla de horas de un curso, con las mismas asignaturas que el Excel 2026."""
        malla = DatosColegio.get_malla_curricular()[curso]
        lecciones = {a: (h, DatosColegio.docentes_de(d)) for a, h, d, _ in malla}
        total = sum(h for _, h, _, _ in malla)

        def horas(*asigs):
            return sum(lecciones[a][0] for a in asigs if a in lecciones)

        def profes(*asigs):
            docs = []
            for a in asigs:
                for d in lecciones.get(a, (0, ()))[1]:
                    if d not in docs:
                        docs.append(d)
            return ' / '.join(docs)

        def fila(etiqueta, *asigs):
            return ('fila', etiqueta, horas(*asigs), profes(*asigs))

        encabezado = ('encabezado', 'ASIGNATURAS', None, None)
        if DatosColegio.es_parvulo(curso):
            return [encabezado, fila('EDUCACIÓN FÍSICA', 'Educación Física y Salud'), ('total', None, total, None)]

        ingles = profes('Idioma Extranjero Inglés')
        if 'English Skills' in lecciones:
            ingles += f" / {lecciones['English Skills'][1][1]} (E. skills)"
        fila_ingles = ('fila', 'IDIOMA EXTRANJERO: INGLÉS', horas('Idioma Extranjero Inglés', 'English Skills'), ingles)
        numero = int(curso[0])

        if 'BÁSICO' in curso and numero <= 4:
            orientec = profes('Orientación / Tecnología')
            media_hora = horas('Orientación / Tecnología') / 2
            filas = [
                encabezado,
                fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                fila_ingles,
                fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Ciencias Sociales'),
                fila('CIENCIAS NATURALES', 'Ciencias Naturales'),
                ('fila', 'EDUCACIÓN TECNOLÓGICA', media_hora, orientec),
                fila('ARTES VISUALES', 'Artes Visuales'),
                fila('MÚSICA', 'Educación Musical'),
                fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                ('fila', 'ORIENTACIÓN', media_hora, orientec),
                fila('RELIGIÓN', 'Religión'),
            ]
        elif 'BÁSICO' in curso:
            filas = [
                encabezado,
                fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                fila_ingles,
                fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Historia, Geografía y CC.SS.'),
                fila('CIENCIAS NATURALES', 'Ciencias Naturales'),
                fila('EDUCACIÓN TECNOLÓGICA', 'Educación Tecnológica'),
            ]
            if numero >= 7:
                filas.append(fila('FÍSICA', 'Física'))
            filas += [
                fila('ARTES VISUALES', 'Artes Visuales'),
                fila('MÚSICA', 'Educación Musical'),
                fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                fila('ORIENTACIÓN', 'Orientación'),
                fila('RELIGIÓN', 'Religión'),
            ]
        else:
            artes, musica = lecciones['Artes Visuales / Música'][1]
            artes_musica = [('artes', 'ARTES VISUALES', horas('Artes Visuales / Música'), artes),
                            ('musica', 'MÚSICA', None, musica)]
            cierre = [fila('EDUCACIÓN FISICA Y SALUD', 'Educación Física y Salud'),
                      fila('ORIENTACIÓN', 'Orientación'),
                      fila('RELIGIÓN', 'Religión')]
            if numero <= 2:
                filas = [
                    encabezado,
                    fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                    fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                    fila_ingles,
                    fila('HISTORIA, GEOGRAFÍA Y CIENCIAS SOCIALES', 'Historia, Geografía y CC.SS.'),
                    ('subtitulo', 'CIENCIAS NATURALES:', None, None),
                    fila('BIOLOGÍA', 'Biología'),
                    fila('QUÍMICA', 'Química'),
                    fila('FÍSICA', 'Física'),
                    fila('EDUCACIÓN TECNOLÓGICA', 'Educación Tecnológica'),
                ] + artes_musica + cierre
            else:
                nivel = DatosColegio.nivel(curso)
                filas = [
                    ('encabezado', 'ASIGNATURAS - FORM. COMUN - LIBRE DISP.', None, None),
                    fila('LENGUAJE Y COMUNICACIÓN', 'Lenguaje y Comunicación'),
                    fila('EDUCACIÓN MATEMÁTICA', 'Educación Matemática'),
                    fila_ingles,
                    fila('EDUCACIÓN CIUDADANA', 'Educación Ciudadana'),
                    fila('CIENCIAS PARA LA CIUDADANIA', 'Ciencias para la Ciudadanía'),
                    fila('FILOSOFÍA', 'Filosofía'),
                    fila('FÍSICA', 'Física'),
                ] + artes_musica + cierre
                filas.append(('encabezado', 'ASIGNATURAS - FORMACIÓN DIFERENCIADA', None, None))
                for electivo in DatosColegio.ELECTIVOS_OFERTA[nivel]:
                    filas.append(('fila', electivo, 6, cls._docentes_electivo(nivel, electivo)))
        filas.append(('total', None, total, None))
        return filas

    @classmethod
    def _hoja_gimnasios(cls, ws, horario):
        cls._formato_hoja(ws)
        recintos = horario.asignar_recintos()
        for rango in ('D1:E1', 'F1:H1', 'F2:H2', 'D3:E3', 'F3:H3', 'D19:E19', 'F18:H18', 'F19:H19',
                      'D34:E34', 'F33:H33', 'F34:H34'):
            ws.merge_cells(rango)
        for recinto, fila_titulo, fila_grilla in (('GIMNASIO A', 2, 5), ('GIMNASIO B', 18, 20),
                                                  ('PATIO SANTO DOMINGO', 33, 35)):
            cls._combinar(ws, fila_titulo, 4, fila_titulo, 5, recinto, negrita=True, tam=10, borde=False, v='center')

            def contenido(dia, b, recinto=recinto):
                if b == 'C':
                    return None, None
                cursos = recintos.get((recinto, dia, b))
                if cursos:
                    return ' / '.join(DatosColegio.codigo_corto(c) for c in cursos), None
                # Los pares libres de los gimnasios se marcan "GIMNASIO / DISPONIBLE", como en el Excel 2026
                b1 = b if b % 2 else b - 1
                if recinto != 'PATIO SANTO DOMINGO' and not recintos.get((recinto, dia, b1)) \
                        and not recintos.get((recinto, dia, b1 + 1)):
                    return ('GIMNASIO' if b % 2 else 'DISPONIBLE'), None
                return None, None

            cls._grilla(ws, fila_grilla, 2, 8, contenido)
        for fila, alto in {4: 5.25, 17: 8.25, 19: 1.5, 32: 6.75, 34: 2.25}.items():
            ws.row_dimensions[fila].height = alto

    @classmethod
    def _hoja_salas(cls, ws, horario):
        anchos = {'A': 0.66, 'I': 10.66, 'Q': 10.66}
        for base in (2, 10, 18):
            anchos[get_column_letter(base)] = 8.66
            anchos[get_column_letter(base + 1)] = 12.66
            for i in range(2, 7):
                anchos[get_column_letter(base + i)] = 13.89
        cls._formato_hoja(ws, anchos)

        contenido_salas = defaultdict(list)   # (sala, dia, bloque) -> [textos]
        for curso in DatosColegio.CURSOS:
            for dia in DatosColegio.DIAS:
                for b, item in horario.asignaciones[curso][dia].items():
                    if item['asignatura'] == 'English Skills':
                        contenido_salas[('ELECTIVO 1', dia, b)].append(DatosColegio.codigo_corto(curso))
                    elif item['asignatura'] == 'Artes Visuales / Música':
                        contenido_salas[('SALA DE ARTES', dia, b)].append(DatosColegio.codigo_sala(curso))
                        contenido_salas[('SALA DE MÚSICA', dia, b)].append(DatosColegio.codigo_sala(curso))
                    elif item['espacio'] == 'Sala de Música':
                        contenido_salas[('SALA DE MÚSICA', dia, b)].append(DatosColegio.codigo_sala(curso))
        for franjas in DatosColegio.FRANJAS_ELECTIVOS.values():
            for franja in franjas:
                for dia, b0 in franja['periodos']:
                    for b in (b0, b0 + 1):
                        for _, _, _, sala, etiqueta in franja['grupos']:
                            contenido_salas[(sala, dia, b)].append(etiqueta)

        for fila_titulo, col0, sala, subtitulo in cls.DISPOSICION_SALAS:
            cls._combinar(ws, fila_titulo, col0 + 2, fila_titulo, col0 + 3, sala, negrita=True, tam=11, borde=False)
            cls._combinar(ws, fila_titulo, col0 + 4, fila_titulo, col0 + 6, subtitulo, negrita=True, borde=False)
            ws.row_dimensions[fila_titulo + 2].height = 5.25
            bloqueos = DatosColegio.BLOQUEOS_PASTORAL.get(sala, [])

            def contenido(dia, b, sala=sala, bloqueos=bloqueos):
                if (dia, b) in bloqueos:
                    return 'PASTORAL', cls.PASTORAL
                textos = contenido_salas.get((sala, dia, b))
                return (' / '.join(textos) if textos else None), None

            cls._grilla(ws, fila_titulo + 3, col0, 10, contenido)

    @classmethod
    def _hoja_auditoria(cls, ws, horario):
        auditoria = ValidadorRestricciones.auditar(horario)
        m = auditoria['metricas']
        cls._formato_hoja(ws, {'A': 0.66, 'B': 46, 'C': 42, 'D': 22})
        cls._escribir(ws, 2, 2, f'INFORME DE CUMPLIMIENTO Y FACTIBILIDAD — AÑO ESCOLAR {DatosColegio.ANIO}',
                      negrita=True, tam=11, borde=False, h=None)
        estado = ('FACTIBLE (100% restricciones duras cumplidas)' if auditoria['valido']
                  else f"NO FACTIBLE ({len(auditoria['errores_duros'])} errores)")
        cls._escribir(ws, 3, 2, 'Estado global:', negrita=True, borde=False, h=None)
        cls._escribir(ws, 3, 3, estado, negrita=True, borde=False, h=None)

        filas = [
            ('Colisiones docentes (no solapamiento)', f"{m['colisiones_docentes']} colisiones", m['colisiones_docentes']),
            ('Cargas curriculares y jornadas completas',
             f"{m['total_bloques_asignados']} / {m['total_bloques_esperados']} bloques", m['errores_carga']),
            ('Orientación a cargo del profesor jefe', f"{m['errores_jefatura']} inconsistencias", m['errores_jefatura']),
            ('Sincronización de electivos 3° y 4° medio', f"{m['errores_electivos']} diferencias A/B", m['errores_electivos']),
            ('Recintos deportivos (1 curso por recinto)', f"{m['errores_recintos']} excesos", m['errores_recintos']),
            ('Sala English skills y bloqueos de pastoral', f"{m['errores_salas']} conflictos", m['errores_salas']),
            ('Una sesión diaria por asignatura', f"{m['mismo_dia_violaciones']} repeticiones", m['mismo_dia_violaciones']),
            ('Troncales en bloques consecutivos', f"{m['core_no_consecutivo']} bloques separados", m['core_no_consecutivo']),
            ('Bloques dobles (90 min)', f"{m['porcentaje_bloques_dobles']}% de las horas", None),
            ('Ventanas docentes', f"{m['total_ventanas_docentes']} bloques de espera", None),
            ('Uso del Patio Santo Domingo', f"{m['uso_patio']} bloques", None),
        ]
        for i, texto in enumerate(['MÉTRICA / CRITERIO', 'VALOR EVALUADO', 'ESTADO']):
            cls._escribir(ws, 5, 2 + i, texto, negrita=True, relleno=cls.AMARILLO)
        fila = 6
        for criterio, valor, errores in filas:
            estado = 'INDICADOR' if errores is None else ('CUMPLE' if errores == 0 else 'NO CUMPLE')
            cls._escribir(ws, fila, 2, criterio, negrita=True, relleno=cls.CREMA, h=None)
            cls._escribir(ws, fila, 3, valor, relleno=cls.BLANCO)
            cls._escribir(ws, fila, 4, estado, relleno=cls.BLANCO)
            fila += 1
        if auditoria['errores_duros']:
            fila += 1
            cls._escribir(ws, fila, 2, 'DETALLE DE ERRORES', negrita=True, relleno=cls.AMARILLO, h=None)
            for error in auditoria['errores_duros'][:300]:
                fila += 1
                cls._combinar(ws, fila, 2, fila, 4, error, relleno=cls.BLANCO, h=None)

    # ------------------------------------------------------------------
    # Libro de docentes
    # ------------------------------------------------------------------
    @classmethod
    def exportar_horario_docentes(cls, horario, ruta_salida):
        """
        Crea el libro de docentes:
        - Hoja 'DISTRIBUCIÓN HORARIA' con el formato de DISTRIBUCIÓN HORARIA 2026.docx
          (una tabla por departamento: profesor, cursos, n° de horas, asignatura y horas lectivas).
        - Una hoja por docente con su grilla semanal y el mismo diseño de las hojas de curso.
        """
        wb = cls._nuevo_libro()
        os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
        ws = wb.active
        ws.title = 'DISTRIBUCIÓN HORARIA'
        cls._hoja_distribucion(ws, horario)
        for depto, docentes in DatosColegio.DEPARTAMENTOS:
            for doc in docentes:
                cls._hoja_docente(wb.create_sheet(cls._nombre_hoja(doc)), doc, depto, horario)
        wb.save(ruta_salida)
        return ruta_salida

    @classmethod
    def _resumen_docente(cls, doc, horario):
        """Filas de la distribución horaria de un docente, calculadas desde el horario generado."""
        orden = {c: i for i, c in enumerate(DatosColegio.todos_los_cursos())}
        normales = defaultdict(int)
        electivos = defaultdict(list)
        for (dia, b), (curso, asig) in horario.docente_ocupado.get(doc, {}).items():
            if curso in orden:
                normales[(curso, asig)] += 1
            else:
                nombre, seccion = asig.rsplit(' (', 1)
                electivos[(curso, nombre)].append(seccion.rstrip(')'))

        filas = []
        for curso, asig in sorted(normales, key=lambda k: (k[1] == 'Orientación', orden[k[0]])):
            if asig == 'Orientación':
                nombre, tipo = 'Jefatura', 'jefatura'
            elif asig == 'English Skills':
                nombre, tipo = 'Inglés (English skills)', 'skills'
            elif asig == 'Orientación / Tecnología':
                nombre, tipo = 'Tecn/C.Orien.', 'orientec'
            elif asig == 'Artes Visuales / Música':
                nombre, tipo = ('Música' if doc == 'Música' else 'Artes'), 'normal'
            else:
                nombre, tipo = cls.NOMBRES_DISTRIBUCION.get(asig, asig), 'normal'
            filas.append({'cursos': DatosColegio.codigo_corto(curso), 'curso': curso,
                          'horas': normales[(curso, asig)], 'asignatura': nombre, 'tipo': tipo})
        for (grupo, nombre), secciones in electivos.items():
            unicas = list(dict.fromkeys(secciones))
            cursos = ', '.join(unicas)
            if all(s.startswith('S.') for s in unicas):   # Secciones mixtas A+B: "3°AB S.1, S.2"
                cursos = f"{grupo.replace(' MEDIO A-B', 'AB')} {cursos}"
            posicion = sum(1 for f in filas if f['tipo'] != 'jefatura')
            filas.insert(posicion, {'cursos': cursos,
                                    'curso': f"{grupo} ({', '.join(unicas)})",
                                    'horas': len(secciones), 'asignatura': nombre, 'tipo': 'electivo'})
        for cursos, actividad, horas in DatosColegio.HORAS_NO_LECTIVAS.get(doc, []):
            filas.append({'cursos': cursos, 'curso': cursos or 'No programada en grilla', 'horas': horas,
                          'asignatura': actividad, 'tipo': 'no_lectiva'})
        return filas, sum(f['horas'] for f in filas)

    @classmethod
    def _hoja_distribucion(cls, ws, horario):
        for col, ancho in {'A': 26, 'B': 20, 'C': 13, 'D': 34, 'E': 16}.items():
            ws.column_dimensions[col].width = ancho
        ws.page_margins = PageMargins(left=0.63, right=0.63, top=0.94, bottom=0.94)
        cls._combinar(ws, 1, 1, 1, 5, f'DISTRIBUCIÓN DE HORAS POR ASIGNATURAS – AÑO ACADÉMICO {DatosColegio.ANIO}.',
                      negrita=True, tam=14, borde=False, h=None, fuente='Calibri')
        rellenos = {'jefatura': cls.CELESTE, 'skills': cls.VERDE, 'orientec': cls.AZUL}
        estilo = {'tam': 10, 'fuente': 'Abadi', 'v': 'center', 'color_borde': cls.BORDE_DOCUMENTO}
        fila = 3
        for titulo, docentes in DatosColegio.DEPARTAMENTOS:
            cls._combinar(ws, fila, 1, fila, 5, f'ASIGNATURAS: {titulo}', negrita=True, tam=12, borde=False,
                          h=None, ajuste=True, fuente='Calibri')
            ws.row_dimensions[fila].height = 32 if len(titulo) > 60 else 18
            fila += 1
            if titulo == 'INGLÉS':
                cls._combinar(ws, fila, 1, fila, 5, 'Verdes: English skills (2 docentes en el mismo bloque)',
                              tam=10, borde=False, h=None, fuente='Calibri')
                fila += 1
            for i, texto in enumerate(['PROFESOR', 'CURSOS', 'Nº DE HORAS', 'ASIGNATURA', 'HORAS LECTIVAS'], 1):
                cls._escribir(ws, fila, i, texto, negrita=True, relleno=cls.GRIS, ajuste=True, **estilo)
            fila += 1
            for doc in docentes:
                filas, total = cls._resumen_docente(doc, horario)
                if not filas:
                    continue
                inicio = fila
                for f in filas:
                    relleno = rellenos.get(f['tipo'])
                    cls._escribir(ws, fila, 2, f['cursos'], relleno=relleno, **estilo)
                    cls._escribir(ws, fila, 3, f['horas'], relleno=relleno, **estilo)
                    cls._escribir(ws, fila, 4, f['asignatura'], h=None, **estilo)
                    fila += 1
                cls._combinar(ws, inicio, 1, fila - 1, 1, doc, **estilo)
                cls._combinar(ws, inicio, 5, fila - 1, 5, total, **estilo)
            fila += 1

    @classmethod
    def _hoja_docente(cls, ws, doc, depto, horario):
        anchos = dict(cls.ANCHOS)
        anchos.update({col: 17 for col in 'DEFGH'})
        cls._formato_hoja(ws, anchos)
        cls._insertar_escudo(ws)
        ocupado = horario.docente_ocupado.get(doc, {})
        detalle = horario.docente_detalle.get(doc, {})
        filas, total = cls._resumen_docente(doc, horario)
        horas = total if total == len(ocupado) else f"{total} ({len(ocupado)} en grilla)"
        departamento = depto.split(' –')[0].split(' - ')[0]
        cls._encabezado(ws, [('DOCENTE', doc), ('DEPARTAMENTO', departamento), ('HORAS LECTIVAS', horas)])

        def contenido(dia, b):
            if b == 'C' or (dia, b) not in ocupado:
                return None, None
            if (dia, b) in detalle:
                return detalle[(dia, b)]['etiqueta'], None
            curso, asig = ocupado[(dia, b)]
            etiqueta = horario.asignaciones[curso][dia][b]['etiqueta']
            if asig == 'Artes Visuales / Música':
                etiqueta = 'MÚSICA' if doc == 'Música' else 'ARTES'
            return f"{DatosColegio.codigo_corto(curso)} {etiqueta}", (cls.SKILLS if asig == 'English Skills' else None)

        max_b = 10 if any(b > 8 for (_, b) in ocupado) else 8
        ultima = cls._grilla(ws, 8, 2, max_b, contenido)
        tabla = [('encabezado', 'ASIGNATURAS', None, None)]
        for f in filas:
            tabla.append(('fila', f['asignatura'].upper(), f['horas'], f['curso']))
        tabla.append(('total', None, total, None))
        fila_total = cls._tabla_horas(ws, ultima + (3 if max_b == 8 else 2), tabla, rotulo_tercera='CURSO')
        if any(f['tipo'] == 'skills' for f in filas):
            cls._leyenda_skills(ws, fila_total + 2)
